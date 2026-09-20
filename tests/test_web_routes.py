"""Pruebas de extremo a extremo sobre la web: quién puede entrar dónde."""
import os

import pytest
from fastapi.testclient import TestClient

from thegrill import db
from thegrill.models import Record, Restaurant, Role, User
from thegrill.web import app as webapp


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("GRILL_INSECURE_COOKIE", "1")     # TestClient habla http
    monkeypatch.setattr(webapp, "UPLOAD_DIR", str(tmp_path / "uploads"))
    db.init_engine(f"sqlite:///{tmp_path/'w.db'}")
    db.create_all()
    with TestClient(webapp.app, follow_redirects=False) as c:
        yield c


def signup(client, restaurant="Casa Pepe", email="ana@casa.com", name="Ana"):
    r = client.post("/signup", data={"restaurant": restaurant, "name": name,
                                     "email": email, "password": "clave-larga-1"})
    assert r.status_code == 303 and r.headers["location"] == "/manager"
    return r.cookies.get(webapp.auth.COOKIE_NAME)


def join(client, code, email, name="Luis"):
    r = client.post("/join", data={"join_code": code, "name": name, "email": email,
                                   "password": "clave-larga-2"})
    assert r.status_code == 303 and r.headers["location"] == "/app"
    return r.cookies.get(webapp.auth.COOKIE_NAME)


def join_code(slug="casa-pepe"):
    with db.session_scope() as s:
        return s.query(Restaurant).filter_by(slug=slug).one().join_code


def csrf_from(html):
    import re
    m = re.search(r'name="csrf" value="([^"]+)"', html)
    assert m, "la página no trae token CSRF"
    return m.group(1)


# ------------------------------------------------------------- sin sesión
def test_anonymous_is_sent_to_login(client):
    for path in ("/", "/app", "/manager", "/manager/equipo", "/app/mis-registros"):
        r = client.get(path)
        assert r.status_code == 303 and r.headers["location"] == "/login", path


def test_health_is_public(client):
    assert client.get("/healthz").json() == {"status": "ok"}


def test_bad_login_does_not_reveal_whether_the_email_exists(client):
    signup(client)
    client.cookies.clear()
    a = client.post("/login", data={"email": "ana@casa.com", "password": "mala"})
    b = client.post("/login", data={"email": "nadie@nada.com", "password": "mala"})
    assert "Email o contraseña incorrectos" in a.text
    assert a.text == b.text


# -------------------------------------------------------------- empleado
def test_employee_sees_capture_screen_but_not_the_manager_area(client):
    signup(client)
    code = join_code()
    client.cookies.clear()
    join(client, code, "luis@casa.com")

    home = client.get("/app")
    assert home.status_code == 200
    assert "¿Qué vas a registrar?" in home.text
    assert "Temperatura de refrigeración" in home.text and "Temperatura de congelación" in home.text

    for path in ("/manager", "/manager/registros", "/manager/alertas",
                 "/manager/plantillas", "/manager/equipo", "/manager/export.csv"):
        r = client.get(path)
        assert r.status_code == 403, path
        assert "solo para managers" in r.text


def test_employee_cannot_create_templates_or_change_roles(client):
    signup(client)
    code = join_code()
    client.cookies.clear()
    join(client, code, "luis@casa.com")
    assert client.post("/manager/plantillas/nueva",
                       data={"name": "Mía", "csrf": "x"}).status_code == 403
    assert client.post("/manager/equipo/1/rol", data={"role": "MANAGER", "csrf": "x"}).status_code == 403


def test_employee_submits_a_record_with_a_photo(client):
    signup(client)
    code = join_code()
    client.cookies.clear()
    join(client, code, "luis@casa.com")

    form = client.get("/app/registro/merma")
    assert form.status_code == 200
    token = csrf_from(form.text)

    sent = client.post("/app/registro/merma",
                       data={"csrf": token, "producto": "Pollo", "cantidad": "2,5",
                             "motivo": "Caducado", "area": "Cocina caliente",
                             "note": "bandeja del fondo"},
                       files={"photos": ("merma.jpg", b"\xff\xd8foto", "image/jpeg")})
    assert sent.status_code == 200 and "Registro guardado" in sent.text

    with db.session_scope() as s:
        rec = s.query(Record).one()
        assert rec.note == "bandeja del fondo"
        assert {v.field_key: v.display for v in rec.values}["cantidad"] == "2.5"   # coma decimal aceptada
        assert len(rec.attachments) == 1

    mine = client.get("/app/mis-registros")
    assert "Merma y desperdicio" in mine.text


def test_form_errors_come_back_on_the_page_without_saving(client):
    signup(client)
    code = join_code()
    client.cookies.clear()
    join(client, code, "luis@casa.com")
    token = csrf_from(client.get("/app/registro/merma").text)
    r = client.post("/app/registro/merma", data={"csrf": token, "cantidad": "2"})
    assert r.status_code == 200 and "Revisa los campos marcados" in r.text
    assert "es obligatorio" in r.text
    with db.session_scope() as s:
        assert s.query(Record).count() == 0


def test_post_without_csrf_is_rejected(client):
    signup(client)
    code = join_code()
    client.cookies.clear()
    join(client, code, "luis@casa.com")
    r = client.post("/app/registro/merma",
                    data={"producto": "Pollo", "cantidad": "1", "motivo": "Rotura", "area": "Almacén"})
    assert r.status_code == 403
    with db.session_scope() as s:
        assert s.query(Record).count() == 0


def test_unknown_template_is_a_clean_404(client):
    signup(client)
    assert client.get("/app/registro/no-existe").status_code == 404


# --------------------------------------------------------------- manager
def test_manager_dashboard_shows_stats_and_join_code(client):
    signup(client)
    code = join_code()
    manager_cookie = client.cookies.get(webapp.auth.COOKIE_NAME)
    client.cookies.clear()
    join(client, code, "luis@casa.com")
    token = csrf_from(client.get("/app/registro/temp_refrigeracion").text)
    client.post("/app/registro/temp_refrigeracion",
                data={"csrf": token, "unidad": "Cámara refrigerada 1", "temperatura": "11"})

    client.cookies.set(webapp.auth.COOKIE_NAME, manager_cookie)
    panel = client.get("/manager")
    assert panel.status_code == 200
    assert code in panel.text                       # el manager reparte el código
    assert "Cumplimiento" in panel.text and "Alertas sin cerrar" in panel.text
    assert "por encima del máximo 5°C" in panel.text

    registros = client.get("/manager/registros")
    assert "Luis" in registros.text                 # ve quién lo registró


def test_manager_closes_an_alert_with_a_corrective_action(client):
    signup(client)
    code = join_code()
    manager_cookie = client.cookies.get(webapp.auth.COOKIE_NAME)
    client.cookies.clear()
    join(client, code, "luis@casa.com")
    token = csrf_from(client.get("/app/registro/temp_congelacion").text)
    client.post("/app/registro/temp_congelacion",
                data={"csrf": token, "unidad": "Congelador 1", "temperatura": "2"})

    client.cookies.set(webapp.auth.COOKIE_NAME, manager_cookie)
    alerts = client.get("/manager/alertas")
    mtoken = csrf_from(alerts.text)
    with db.session_scope() as s:
        from thegrill.models import Alert
        alert_id = s.query(Alert).one().id
    closed = client.post(f"/manager/alertas/{alert_id}/cerrar",
                         data={"csrf": mtoken, "resolution": "Puerta mal cerrada, producto trasladado"})
    assert closed.status_code == 303
    assert "Ninguna alerta abierta" in client.get("/manager").text


def test_manager_creates_a_template_and_the_employee_can_fill_it(client):
    signup(client)
    code = join_code()
    manager_cookie = client.cookies.get(webapp.auth.COOKIE_NAME)
    token = csrf_from(client.get("/manager/plantillas").text)
    r = client.post("/manager/plantillas/nueva",
                    data={"csrf": token, "name": "Aceite de freidora", "category": "haccp",
                          "expected_per_day": "1",
                          "fields_spec": "Freidora | TEXT\nTemperatura | NUMBER | °C | 150 | 180"})
    assert r.status_code == 303

    client.cookies.clear()
    join(client, code, "luis@casa.com")
    form = client.get("/app/registro/aceite_de_freidora")
    assert form.status_code == 200 and "Temperatura" in form.text
    ftoken = csrf_from(form.text)
    sent = client.post("/app/registro/aceite_de_freidora",
                       data={"csrf": ftoken, "freidora": "Freidora 1", "temperatura": "195"})
    assert "Registro guardado con alerta" in sent.text
    assert "por encima del máximo 180°C" in sent.text


def test_manager_can_deactivate_a_template(client):
    signup(client)
    token = csrf_from(client.get("/manager/plantillas").text)
    with db.session_scope() as s:
        from thegrill.models import RecordTemplate
        tpl_id = s.query(RecordTemplate).filter_by(code="inventario").one().id
    assert client.post(f"/manager/plantillas/{tpl_id}/activar", data={"csrf": token}).status_code == 303
    assert client.get("/app/registro/inventario").status_code == 404


def test_manager_promotes_and_deactivates_staff(client):
    signup(client)
    code = join_code()
    manager_cookie = client.cookies.get(webapp.auth.COOKIE_NAME)
    client.cookies.clear()
    join(client, code, "luis@casa.com")
    client.cookies.set(webapp.auth.COOKIE_NAME, manager_cookie)

    team = client.get("/manager/equipo")
    token = csrf_from(team.text)
    with db.session_scope() as s:
        luis_id = s.query(User).filter_by(email="luis@casa.com").one().id
        manager_id = s.query(User).filter_by(email="ana@casa.com").one().id

    assert client.post(f"/manager/equipo/{luis_id}/rol",
                       data={"csrf": token, "role": "MANAGER"}).status_code == 303
    with db.session_scope() as s:
        assert s.get(User, luis_id).role == Role.MANAGER

    assert client.post(f"/manager/equipo/{manager_id}/rol",
                       data={"csrf": token, "role": "EMPLOYEE"}).status_code == 400   # no a sí mismo
    assert client.post(f"/manager/equipo/{manager_id}/activar",
                       data={"csrf": token}).status_code == 400

    assert client.post(f"/manager/equipo/{luis_id}/activar", data={"csrf": token}).status_code == 303
    client.cookies.clear()
    blocked = client.post("/login", data={"email": "luis@casa.com", "password": "clave-larga-2"})
    assert "Cuenta desactivada" in blocked.text


def test_csv_export_downloads(client):
    signup(client)
    token = csrf_from(client.get("/app/registro/limpieza").text)
    client.post("/app/registro/limpieza",
                data={"csrf": token, "zona": "Cocina", "realizada": "1", "producto_usado": "Lejía"})
    r = client.get("/manager/export.csv?days=7")
    assert r.status_code == 200
    assert "attachment; filename=" in r.headers["content-disposition"]
    assert "Limpieza y desinfección" in r.text and "Lejía" in r.text


# ------------------------------------------- aislamiento entre restaurantes
def test_one_restaurant_never_sees_anothers_data(client):
    signup(client, "Casa Pepe", "ana@casa.com", "Ana")
    token = csrf_from(client.get("/app/registro/merma").text)
    client.post("/app/registro/merma", data={"csrf": token, "producto": "Secreto de Casa Pepe",
                                             "cantidad": "1", "motivo": "Rotura", "area": "Almacén"})
    client.cookies.clear()

    signup(client, "El Otro", "eva@otro.com", "Eva")
    assert "Secreto de Casa Pepe" not in client.get("/manager/registros").text
    assert client.get("/manager").text.count("Secreto") == 0
    assert "Secreto de Casa Pepe" not in client.get("/manager/export.csv").text


def test_photos_are_not_served_across_restaurants(client):
    signup(client, "Casa Pepe", "ana@casa.com", "Ana")
    token = csrf_from(client.get("/app/registro/merma").text)
    client.post("/app/registro/merma",
                data={"csrf": token, "producto": "Pollo", "cantidad": "1",
                      "motivo": "Rotura", "area": "Almacén"},
                files={"photos": ("p.jpg", b"\xff\xd8secreto", "image/jpeg")})
    with db.session_scope() as s:
        from thegrill.models import Attachment
        photo_id = s.query(Attachment).one().id
    assert client.get(f"/foto/{photo_id}").status_code == 200

    client.cookies.clear()
    signup(client, "El Otro", "eva@otro.com", "Eva")
    r = client.get(f"/foto/{photo_id}")
    assert r.status_code == 403 and "otro restaurante" in r.text


def test_an_employee_only_sees_his_own_photos(client):
    signup(client, "Casa Pepe", "ana@casa.com", "Ana")
    code = join_code()
    manager_cookie = client.cookies.get(webapp.auth.COOKIE_NAME)
    client.cookies.clear()
    join(client, code, "luis@casa.com", "Luis")
    token = csrf_from(client.get("/app/registro/merma").text)
    client.post("/app/registro/merma",
                data={"csrf": token, "producto": "Pollo", "cantidad": "1",
                      "motivo": "Rotura", "area": "Almacén"},
                files={"photos": ("p.jpg", b"\xff\xd8foto", "image/jpeg")})
    with db.session_scope() as s:
        from thegrill.models import Attachment
        photo_id = s.query(Attachment).one().id
    assert client.get(f"/foto/{photo_id}").status_code == 200       # la suya sí

    client.cookies.clear()
    join(client, code, "otro@casa.com", "Otro")
    assert client.get(f"/foto/{photo_id}").status_code == 403       # la de un compañero no

    client.cookies.set(webapp.auth.COOKIE_NAME, manager_cookie)
    assert client.get(f"/foto/{photo_id}").status_code == 200       # el manager sí


def test_logout_kills_the_session(client):
    signup(client)
    assert client.get("/manager").status_code == 200
    assert client.post("/logout").status_code == 303
    assert client.get("/manager").headers["location"] == "/login"


# ------------------------------------------------------- notificaciones web
def test_the_badge_and_the_feed_reach_the_manager(client):
    signup(client)
    code = join_code()
    manager_cookie = client.cookies.get(webapp.auth.COOKIE_NAME)

    client.cookies.clear()
    join(client, code, "luis@casa.com")
    token = csrf_from(client.get("/app/registro/temp_refrigeracion").text)
    client.post("/app/registro/temp_refrigeracion",
                data={"csrf": token, "unidad": "Vitrina", "temperatura": "11"})
    feed = client.get("/api/notificaciones").json()
    assert feed["unread"] == 0                       # quien lo registró ya lo vio al guardar

    client.cookies.set(webapp.auth.COOKIE_NAME, manager_cookie)
    feed = client.get("/api/notificaciones").json()
    assert feed["unread"] == 1
    item = feed["items"][0]
    assert item["severity"] == "CRITICAL" and "Luis" in item["body"]

    panel = client.get("/manager")
    assert 'id="navbadge"' in panel.text and "hidden" not in panel.text.split('id="navbadge"')[1][:40]

    avisos = client.get("/notificaciones")
    assert "11°C" in avisos.text and "Ver y cerrar la alerta" in avisos.text
    assert client.get("/api/notificaciones").json()["unread"] == 0     # verlos los marca


def test_the_feed_is_private_to_each_user(client):
    signup(client, "Casa Pepe", "ana@casa.com", "Ana")
    code = join_code()
    client.cookies.clear()
    join(client, code, "luis@casa.com")
    token = csrf_from(client.get("/app/registro/temp_refrigeracion").text)
    client.post("/app/registro/temp_refrigeracion",
                data={"csrf": token, "unidad": "Vitrina", "temperatura": "11"})

    client.cookies.clear()
    signup(client, "El Otro", "eva@otro.com", "Eva")
    assert client.get("/api/notificaciones").json() == {"unread": 0, "items": []}
    assert "Vitrina" not in client.get("/notificaciones").text


def test_the_feed_needs_a_session(client):
    r = client.get("/api/notificaciones")
    assert r.status_code == 303 and r.headers["location"] == "/login"


def test_an_employee_is_told_how_his_alert_was_resolved(client):
    signup(client)
    code = join_code()
    manager_cookie = client.cookies.get(webapp.auth.COOKIE_NAME)
    client.cookies.clear()
    employee_cookie = join(client, code, "luis@casa.com")
    token = csrf_from(client.get("/app/registro/temp_refrigeracion").text)
    client.post("/app/registro/temp_refrigeracion",
                data={"csrf": token, "unidad": "Vitrina", "temperatura": "11"})

    client.cookies.set(webapp.auth.COOKIE_NAME, manager_cookie)
    alerts = client.get("/manager/alertas")
    with db.session_scope() as s:
        from thegrill.models import Alert
        alert_id = s.query(Alert).one().id
    client.post(f"/manager/alertas/{alert_id}/cerrar",
                data={"csrf": csrf_from(alerts.text), "resolution": "Vitrina reparada"})

    client.cookies.set(webapp.auth.COOKIE_NAME, employee_cookie)
    assert client.get("/api/notificaciones").json()["unread"] == 1
    assert "Vitrina reparada" in client.get("/notificaciones").text


# --------------------------------------------------------------- idiomas
def test_login_page_offers_every_language(client):
    html = client.get("/login").text
    for name in ("Español", "English", "Français", "Deutsch", "Nederlands", "العربية"):
        assert name in html, name
    assert 'href="/idioma/de?next=/login"' in html


def test_choosing_a_language_on_login_changes_the_page_and_sticks(client):
    assert "Entrar" in client.get("/login").text
    r = client.get("/idioma/de?next=/login")
    assert r.status_code == 303 and r.headers["location"] == "/login"
    assert client.cookies.get(webapp.i18n.COOKIE_NAME) == "de"
    html = client.get("/login").text
    assert "Anmelden" in html and "Zum ersten Mal hier?" in html


def test_arabic_switches_the_page_to_right_to_left(client):
    client.get("/idioma/ar?next=/login")
    html = client.get("/login").text
    assert '<html lang="ar" dir="rtl">' in html
    assert "تسجيل الدخول" in html


def test_an_unknown_language_is_rejected(client):
    assert client.get("/idioma/klingon").status_code == 404
    assert webapp.i18n.COOKIE_NAME not in client.cookies


def test_the_language_redirect_cannot_be_used_to_send_people_elsewhere(client):
    for bad in ("https://evil.example/x", "//evil.example/x"):
        r = client.get(f"/idioma/en?next={bad}")
        assert r.headers["location"] == "/login"


def test_the_browser_language_is_honoured_before_logging_in(client):
    html = client.get("/login", headers={"accept-language": "nl-BE,nl;q=0.9,fr;q=0.8"}).text
    assert "Aanmelden" in html
    html = client.get("/login", headers={"accept-language": "ja,ko"}).text
    assert "Entrar" in html            # ninguno disponible: español


def test_signing_up_in_dutch_creates_dutch_forms(client):
    r = client.post("/signup", data={"restaurant": "De Kombuis", "name": "Piet",
                                     "email": "piet@kombuis.be", "password": "clave-larga-1",
                                     "language": "nl"})
    assert r.status_code == 303
    panel = client.get("/manager").text
    assert "Naleving" in panel and "Open waarschuwingen" in panel
    assert "Koeltemperatuur" in client.get("/app").text


def test_settings_change_only_my_language_not_my_colleagues(client):
    signup(client, "Casa Pepe", "ana@casa.com", "Ana")       # se crea en español
    manager_cookie = client.cookies.get(webapp.auth.COOKIE_NAME)
    code = join_code()

    client.cookies.clear()
    join(client, code, "hans@casa.com", "Hans")
    assert "¿Qué vas a registrar?" in client.get("/app").text
    settings = client.get("/configuracion")
    assert settings.status_code == 200
    saved = client.post("/configuracion", data={"csrf": csrf_from(settings.text), "language": "de"})
    assert saved.status_code == 303
    assert "Was möchtest du erfassen?" in client.get("/app").text

    client.cookies.set(webapp.auth.COOKIE_NAME, manager_cookie)
    assert "Cumplimiento" in client.get("/manager").text     # la manager sigue en español


def test_an_employee_cannot_change_the_restaurant_language(client):
    signup(client, "Casa Pepe", "ana@casa.com", "Ana")
    code = join_code()
    client.cookies.clear()
    join(client, code, "luis@casa.com")
    settings = client.get("/configuracion")
    assert "settings.restaurant_language" not in settings.text
    assert 'name="restaurant_language"' not in settings.text     # ni siquiera se le ofrece
    client.post("/configuracion", data={"csrf": csrf_from(settings.text),
                                        "language": "fr", "restaurant_language": "fr"})
    with db.session_scope() as s:
        from thegrill.models import Restaurant
        assert s.query(Restaurant).filter_by(slug="casa-pepe").one().language == "es"


def test_a_manager_can_change_the_restaurant_language(client):
    signup(client, "Casa Pepe", "ana@casa.com", "Ana")
    settings = client.get("/configuracion")
    client.post("/configuracion", data={"csrf": csrf_from(settings.text),
                                        "language": "fr", "restaurant_language": "fr"})
    with db.session_scope() as s:
        from thegrill.models import Restaurant
        assert s.query(Restaurant).filter_by(slug="casa-pepe").one().language == "fr"
    assert "Conformité" in client.get("/manager").text


def test_settings_need_a_csrf_token(client):
    signup(client)
    assert client.post("/configuracion", data={"language": "de"}).status_code == 403


def test_alerts_stay_in_the_restaurant_language_whatever_the_reader_uses(client):
    signup(client, "Casa Pepe", "ana@casa.com", "Ana")       # restaurante en español
    code = join_code()
    manager_cookie = client.cookies.get(webapp.auth.COOKIE_NAME)

    client.cookies.clear()
    join(client, code, "hans@casa.com", "Hans")
    settings = client.get("/configuracion")
    client.post("/configuracion", data={"csrf": csrf_from(settings.text), "language": "de"})
    form = client.get("/app/registro/temp_refrigeracion")
    assert "Eintrag speichern" in form.text                  # la pantalla, en alemán
    sent = client.post("/app/registro/temp_refrigeracion",
                       data={"csrf": csrf_from(form.text), "unidad": "Vitrina", "temperatura": "11"})
    assert "Eintrag mit Warnung gespeichert" in sent.text    # el mensaje de pantalla, en alemán
    assert "por encima del máximo 5°C" in sent.text          # la alerta guardada, en español

    client.cookies.set(webapp.auth.COOKIE_NAME, manager_cookie)
    assert "por encima del máximo 5°C" in client.get("/manager").text


def test_no_screen_leaks_spanish_when_the_language_is_english(client):
    """Barrido: ninguna pantalla se deja texto sin traducir."""
    client.post("/signup", data={"restaurant": "The Kitchen", "name": "Ann",
                                 "email": "ann@kitchen.com", "password": "clave-larga-1",
                                 "language": "en"})
    form = client.get("/app/registro/temp_refrigeracion")
    client.post("/app/registro/temp_refrigeracion",
                data={"csrf": csrf_from(form.text), "unidad": "Display unit", "temperatura": "11"})

    spanish = ["Cumplimiento", "Registrar", "Alertas", "Plantillas", "Equipo", "Avisos",
               "Configuración", "Salir", "Guardar", "Fecha", "Estado", "Descargar",
               "¿Qué vas a registrar?", "Acción correctiva", "Último acceso", "Motivo"]
    pages = ["/manager", "/manager/registros", "/manager/alertas", "/manager/plantillas",
             "/manager/equipo", "/app", "/app/mis-registros", "/notificaciones",
             "/configuracion", "/app/registro/temp_refrigeracion", "/merma"]
    for path in pages:
        html = client.get(path).text
        leaked = [w for w in spanish if w in html]
        assert not leaked, f"{path} deja en español: {leaked}"


# ------------------------------------------------------------- descargas
def test_the_downloads_page_lists_every_active_form(client):
    signup(client)
    html = client.get("/descargas").text
    assert "Hojas para imprimir" in html
    assert "Temperatura de refrigeración" in html and "Conteo de inventario" in html
    assert 'href="/descargas/todo.xlsx"' in html
    assert 'href="/descargas/temp_refrigeracion.xlsx"' in html
    assert "no genera alertas" in html        # el papel hay que pasarlo luego


def test_staff_can_download_the_sheets_too(client):
    signup(client)
    code = join_code()
    client.cookies.clear()
    join(client, code, "luis@casa.com")
    assert client.get("/descargas").status_code == 200
    assert client.get("/descargas/merma.xlsx").status_code == 200


def test_downloading_one_form_gives_a_real_workbook(client):
    import io
    from openpyxl import load_workbook
    signup(client)
    r = client.get("/descargas/temp_refrigeracion.xlsx")
    assert r.status_code == 200
    assert r.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert 'filename="temp_refrigeracion.xlsx"' in r.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Temperatura de refrigeración"]
    assert wb.active["A1"].value == "Temperatura de refrigeración"


def test_downloading_everything_gives_one_tab_per_form(client):
    import io
    from openpyxl import load_workbook
    signup(client)
    r = client.get("/descargas/todo.xlsx")
    wb = load_workbook(io.BytesIO(r.content))
    assert len(wb.sheetnames) == 7
    assert "casa-pepe-hojas.xlsx" in r.headers["content-disposition"]


def test_a_deactivated_form_is_not_downloadable(client):
    signup(client)
    token = csrf_from(client.get("/manager/plantillas").text)
    with db.session_scope() as s:
        from thegrill.models import RecordTemplate
        tpl_id = s.query(RecordTemplate).filter_by(code="inventario").one().id
    client.post(f"/manager/plantillas/{tpl_id}/activar", data={"csrf": token})
    assert client.get("/descargas/inventario.xlsx").status_code == 404
    assert "Conteo de inventario" not in client.get("/descargas").text


def test_you_cannot_download_another_restaurants_sheets(client):
    import io
    from openpyxl import load_workbook
    signup(client, "Casa Pepe", "ana@casa.com", "Ana")
    token = csrf_from(client.get("/manager/plantillas").text)
    client.post("/manager/plantillas/nueva",
                data={"csrf": token, "name": "Secreto de Casa Pepe", "category": "other",
                      "fields_spec": "Dato | TEXT"})
    client.cookies.clear()

    signup(client, "El Otro", "eva@otro.com", "Eva")
    assert client.get("/descargas/secreto_de_casa_pepe.xlsx").status_code == 404
    wb = load_workbook(io.BytesIO(client.get("/descargas/todo.xlsx").content))
    assert "Secreto de Casa Pepe" not in wb.sheetnames


def test_downloads_need_a_session(client):
    for path in ("/descargas", "/descargas/todo.xlsx", "/descargas/merma.xlsx"):
        r = client.get(path)
        assert r.status_code == 303 and r.headers["location"] == "/login", path


def test_the_sheet_follows_the_readers_language(client):
    import io
    from openpyxl import load_workbook
    signup(client)                                  # restaurante en español
    settings = client.get("/configuracion")
    client.post("/configuracion", data={"csrf": csrf_from(settings.text), "language": "nl"})
    assert "Formulieren om af te drukken" in client.get("/descargas").text
    ws = load_workbook(io.BytesIO(client.get("/descargas/merma.xlsx").content)).active
    assert [c.value for c in ws[8]][:3] == ["Nº", "Datum", "Tijd"]
    assert ws.cell(row=9, column=1).value == "VOORBEELD"


# ------------------------------------------------------------- merma
def stocked(kg=5.1, unit_cost=30.0, serial="8017-01", tg="TG-0010"):
    """Un lote de filetes ya cortados, listo para tirar parte de él."""
    from datetime import date, timedelta

    from thegrill.models import Ingredient, IngredientItem, Rotation, Unit
    from thegrill.web import costing
    with db.session_scope() as s:
        rest = s.query(Restaurant).filter_by(slug="casa-pepe").one()
        ana = s.query(User).filter_by(restaurant_id=rest.id, role=Role.MANAGER).first()
        ing = Ingredient(restaurant_id=rest.id, name="Striploin steak",
                         unit=Unit.KG, rotation=Rotation.FEFO)
        s.add(ing); s.flush()
        item = IngredientItem(restaurant_id=rest.id, ingredient_id=ing.id,
                              name="Striploin steak 300g AUS")
        s.add(item); s.flush()
        lot = costing.receive(s, ana, item, kg, unit_cost,
                              date.today() + timedelta(days=10), lot_code=tg)
        lot.serial = serial
        lot.pieces = 17
        s.flush()
        return ing.id


def test_the_waste_page_is_reachable_from_the_meat_screen(client):
    signup(client)
    assert '/merma' in client.get("/carne").text
    assert client.get("/merma").status_code == 200


def test_recording_waste_writes_it_down_and_raises_the_cost_of_the_rest(client):
    from thegrill.models import IngredientLot, IngredientMovement, MovementKind
    signup(client)
    stocked(kg=5.1, unit_cost=30.0)
    form = client.get("/merma")
    r = client.post("/merma", data={"csrf": csrf_from(form.text), "serial": "8017-01",
                                    "kg": "0,6", "pieces": "2", "reason": "Caducado"})
    assert r.status_code == 200
    assert "TG-0010" in r.text and "2 pz" in r.text and "Caducado" in r.text

    with db.session_scope() as s:
        mov = s.query(IngredientMovement).filter_by(kind=MovementKind.WASTE).one()
        assert mov.qty == -0.6 and mov.cost == 18.0
        lot = s.query(IngredientLot).filter_by(serial="8017-01").one()
        assert lot.qty_remaining == 4.5
        assert lot.unit_cost == 34.0                 # los quince que quedan pagan los diecisiete


def test_waste_bigger_than_what_is_left_is_refused_without_touching_stock(client):
    from thegrill.models import IngredientLot, IngredientMovement, MovementKind
    signup(client)
    stocked(kg=5.1)
    form = client.get("/merma")
    r = client.post("/merma", data={"csrf": csrf_from(form.text), "serial": "8017-01",
                                    "kg": "9"})
    assert r.status_code == 200 and "negativo" in r.text
    with db.session_scope() as s:
        assert s.query(IngredientLot).filter_by(serial="8017-01").one().qty_remaining == 5.1
        assert s.query(IngredientMovement).filter_by(kind=MovementKind.WASTE).count() == 0


def test_waste_needs_a_csrf_token(client):
    signup(client)
    stocked()
    assert client.post("/merma", data={"serial": "8017-01", "kg": "0.5"}).status_code == 403


def test_an_employee_can_record_waste_too(client):
    """Quien tira la pieza es quien está en la cámara, no el manager."""
    from thegrill.models import IngredientMovement, MovementKind
    signup(client)
    ing_id = stocked()
    client.cookies.clear()
    join(client, join_code(), "luis@casa.com")
    form = client.get("/merma")
    r = client.post("/merma", data={"csrf": csrf_from(form.text), "serial": "8017-01",
                                    "kg": "0.3", "pieces": "1"})
    assert r.status_code == 200
    with db.session_scope() as s:
        assert s.query(IngredientMovement).filter_by(kind=MovementKind.WASTE).count() == 1
        assert ing_id


def test_waste_needs_a_session(client):
    r = client.get("/merma")
    assert r.status_code == 303 and r.headers["location"] == "/login"
