"""Que el cliente pueda llevarse sus datos, y sobre todo el día que se va.

Esto se escribe después de encontrar que no se podía. La edición de carne
tenía un `/descargas` lleno de botones de Excel, y las catorce hojas que
bajaban eran **rejillas en blanco** para imprimir y rellenar a mano: ni una
consulta a la base de datos. Un restaurante que dejara de pagar perdía el
acceso el mismo día a sus recepciones, sus despieces, sus pesadas y sus
inventarios —los papeles que le piden en una inspección— sin haber tenido
nunca manera de copiarlos.

Las cuatro cosas que se comprueban aquí son las cuatro que pueden volver a
romperse, y cada una se rompe de una forma distinta:

1. **Que baje con datos dentro.** Un libro con las pestañas y sin filas es el
   fallo original con otra cara.
2. **Que baje con la casa bloqueada.** Es la razón de existir de todo esto, y
   es lo más fácil de deshacer sin querer: basta con que alguien cambie el
   `Depends` por el de siempre, que es el que parece correcto.
3. **Que no se cuele la casa de al lado.** Media docena de estas tablas no
   llevan el restaurante escrito encima —los cortes cuelgan del despiece, las
   líneas del inventario—, así que hay que filtrarlas por su madre. Un olvido
   ahí no da error: manda los datos de otro cliente dentro del fichero.
4. **Que los encabezados estén traducidos.** Salen de las claves de texto del
   propio programa; si mañana se renombra una, la columna no falla: escribe
   `m.rec.serial` y nadie se entera hasta que un cliente abre el fichero.
"""
from datetime import date, timedelta

import pytest
from fastapi.testclient import TestClient

from thegrill import db
from thegrill.meat import app as meatapp
from thegrill.meat import service as meat
from thegrill.models import (Billing, CountItemKind, CountPeriod, Ingredient,
                             IngredientItem,
                             MeatCount, MeatCountLine, Primal, Restaurant, Role,
                             Unit, User)
from tests.meat_helpers import SPANISH, add_user, new_house

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.fixture
def casas(tmp_path, monkeypatch):
    """Dos casas con carne dentro. Dos, porque el fallo caro es mezclarlas."""
    pytest.importorskip("openpyxl")
    monkeypatch.setenv("GRILL_INSECURE_COOKIE", "1")
    db.init_engine(f"sqlite:///{tmp_path/'exportar.db'}")
    db.create_all()
    mia = new_house("Asador Marina", "albano@marina.com")
    suya = new_house("Asador Vecino", "vecina@vecino.com", manager="Nora")
    _mete_carne(mia, "M-900")
    _mete_carne(suya, "V-900")
    return mia, suya


def _mete_carne(casa_id: int, serial: str) -> None:
    """Una casa con trabajo dentro, y del que tiene hijas.

    No basta con una recepción. Las tablas que pueden filtrarse mal son las que
    no llevan el restaurante escrito encima —los primales y los cortes de un
    despiece, las líneas de un inventario—, y esas solo existen si en la casa
    se ha despiezado y se ha contado. Con dos recepciones sueltas, la prueba de
    que no se mezclan las casas pasaría sin haber mirado ninguna de las tres.

    Cada dato lleva el prefijo de su casa —«M-» o «V-»— para poder buscarlo
    luego en el fichero del otro.
    """
    with db.session_scope() as session:
        quien = (session.query(User)
                 .filter_by(restaurant_id=casa_id, role=Role.MANAGER).one())
        meat.receive_primals(session, quien, f"{serial}-L", [meat.PrimalRow(
            serial=serial, kg=12.5, price_kg=18.0, sku="RIBEYE", grade="MB7",
            origin="Australia",
            use_by=date.today() + timedelta(days=30))], lang="es")

        madre = Ingredient(restaurant_id=casa_id, name=f"{serial}-ingrediente",
                           unit=Unit.KG)
        session.add(madre)
        session.flush()
        articulo = IngredientItem(restaurant_id=casa_id, ingredient_id=madre.id,
                                  name=f"{serial}-articulo")
        session.add(articulo)
        session.flush()
        meat.post_butchery(session, quien, f"{serial}-TG", [serial], 12.5,
                           [meat.CutRow(name=f"{serial}-corte", item_id=articulo.id,
                                        pieces=10, grams=400.0)], lang="es")

        cuenta = MeatCount(restaurant_id=casa_id, date=date.today(),
                           period=CountPeriod.SPOT, created_by=quien.id)
        session.add(cuenta)
        session.flush()
        session.add(MeatCountLine(count_id=cuenta.id, kind=CountItemKind.PRIMAL,
                                  serial=serial, label=f"{serial}-linea",
                                  counted_kg=1.0))


def entra(email: str, password: str = "clave-larga-1") -> TestClient:
    cliente = TestClient(meatapp.app, follow_redirects=False, headers=SPANISH)
    respuesta = cliente.post("/login", data={"email": email, "password": password})
    assert respuesta.status_code == 303, respuesta.text[:200]
    return cliente


def _abre(payload: bytes):
    import io

    from openpyxl import load_workbook
    return load_workbook(io.BytesIO(payload))


def _todo_el_texto(wb) -> str:
    return "\n".join(str(c.value) for ws in wb.worksheets
                     for fila in ws.iter_rows() for c in fila if c.value is not None)


def _bloquea(casa_id: int) -> None:
    """La casa deja de pagar. Es el día que importa."""
    with db.session_scope() as session:
        session.get(Restaurant, casa_id).billing = Billing.CANCELLED


# --------------------------------------------------------------- que baje
def test_the_manager_downloads_the_whole_house(casas):
    """Con las pestañas y con las filas: un libro vacío es el mismo fallo."""
    cliente = entra("albano@marina.com")
    r = cliente.get("/descargas/mis-datos.xlsx")
    assert r.status_code == 200, r.text[:300]
    assert r.headers["content-type"].startswith(XLSX)
    assert "Asador-Marina" in r.headers["content-disposition"]
    assert str(date.today()) in r.headers["content-disposition"]

    wb = _abre(r.content)
    assert len(wb.worksheets) == 14, [ws.title for ws in wb.worksheets]
    piezas = wb["Recepción de primales"]
    assert piezas.max_row == 2, "la hoja de piezas bajó sin la pieza dentro"
    assert "M-900" in _todo_el_texto(wb)


# ------------------------------------------------- el día que deja de pagar
def test_a_cancelled_house_can_still_take_its_records(casas):
    """Lo que hace falta que siga funcionando cuando ya no hay nada que cobrar.

    Se comprueba primero que la casa está de verdad parada —`/hoy` la manda a
    `/cuenta`—, porque si no la prueba pasaría igual con la cuenta al día y no
    estaría probando nada.
    """
    mia, _ = casas
    _bloquea(mia)
    cliente = entra("albano@marina.com")

    trabajo = cliente.get("/hoy")
    assert trabajo.status_code == 303 and trabajo.headers["location"] == "/cuenta", (
        "esta prueba solo vale con la casa parada de verdad")

    r = cliente.get("/descargas/mis-datos.xlsx")
    assert r.status_code == 200, "el cliente se quedó sin sus papeles al cancelar"
    assert "M-900" in _todo_el_texto(_abre(r.content))

    # Y el enlace tiene que estar donde va a mirar: la única pantalla que ve.
    cuenta = cliente.get("/cuenta")
    assert "/descargas/mis-datos.xlsx" in cuenta.text


# ------------------------------------------------------- nada de la de al lado
def test_the_workbook_never_carries_the_other_house(casas):
    """Las tablas hijas no llevan la casa encima: se filtran por su madre."""
    for correo, mio, ajeno in (("albano@marina.com", "M-900", "V-900"),
                               ("vecina@vecino.com", "V-900", "M-900")):
        texto = _todo_el_texto(_abre(
            entra(correo).get("/descargas/mis-datos.xlsx").content))
        assert mio in texto
        assert ajeno not in texto, f"{correo} se llevó la carne de la otra casa"


# ------------------------------------------------------------- quién y en qué
def test_only_whoever_runs_the_house_can_take_it(casas):
    """Se abre la mano en cuándo, no en quién: un cocinero no se lleva la casa."""
    cocinero = add_user(None, "paco@marina.com", "Paco", "clave-larga-2",
                        role=Role.BUTCHER, house=casas[0])
    assert cocinero.get("/descargas/mis-datos.xlsx").status_code == 403
    assert TestClient(meatapp.app, follow_redirects=False).get(
        "/descargas/mis-datos.xlsx").status_code == 303   # a entrar


def test_the_column_headings_come_out_in_the_language_of_the_house(casas):
    """Y no en claves: `m.rec.serial` en una columna no rompe nada y se ve fatal."""
    mia, _ = casas
    for idioma, esperado in (("es", "Número de primal"), ("de", "Primalnummer")):
        with db.session_scope() as session:
            session.get(Restaurant, mia).language = idioma
            for persona in session.query(User).filter_by(restaurant_id=mia).all():
                persona.language = idioma
        wb = _abre(entra("albano@marina.com").get("/descargas/mis-datos.xlsx").content)
        cabeceras = [c.value for c in wb.worksheets[0][1]]
        assert esperado in cabeceras, cabeceras
        texto = _todo_el_texto(wb)
        assert "m.rec." not in texto and "common." not in texto, "salió una clave sin traducir"


# ============================== y las fotos, que es lo que vale como prueba
def test_what_the_customer_takes_away_includes_the_label_photos(casas, tmp_path):
    """El libro llevaba filas. La etiqueta es la prueba, y no iba.

    El matadero, el lote y la fecha de sacrificio están escritos en la
    etiqueta, y lo que enseña un restaurante en una inspección es la etiqueta,
    no una casilla de una hoja de cálculo que ha escrito él mismo. Este módulo
    existe porque el cliente tiene que poder cumplir con lo que la ley le exige
    a él; llevándose solo las filas se llevaba media casa, y la que faltaba era
    la mitad que vale delante de un inspector.
    """
    import zipfile

    from thegrill.meat import exportar

    casa, _ = casas
    with db.session_scope() as s:
        carpeta = tmp_path / "subidas"
        carpeta.mkdir(exist_ok=True)
        for pieza in s.query(Primal).filter_by(restaurant_id=casa):
            foto = carpeta / f"{pieza.serial}.jpg"
            foto.write_bytes(b"\xff\xd8\xff" + pieza.serial.encode())
            pieza.photo_ref = str(foto)
        seriales = [p.serial for p in s.query(Primal).filter(
            Primal.restaurant_id == casa, Primal.photo_ref.isnot(None))]

    destino = tmp_path / "todo.zip"
    with db.session_scope() as s:
        cuenta = exportar.paquete(s, casa, "es", destino)
    assert seriales, "la casa de la prueba salió sin piezas"
    assert cuenta["fotos"] == len(seriales)

    with zipfile.ZipFile(destino) as z:
        dentro = z.namelist()
        assert exportar.LIBRO in dentro
        for serial in seriales:
            assert f"{exportar.FOTOS}/{serial}.jpg" in dentro, dentro
            assert z.read(f"{exportar.FOTOS}/{serial}.jpg").endswith(serial.encode())


def test_a_photo_that_is_no_longer_on_disk_does_not_cost_him_the_rest(casas, tmp_path):
    """Una foto que falta no puede dejar al cliente sin sus datos.

    Y tampoco puede desaparecer sin decirlo: el que abre el paquete tiene que
    saber que falta y cuántas.
    """
    import zipfile

    from thegrill.meat import exportar

    casa, _ = casas
    with db.session_scope() as s:
        s.query(Primal).filter_by(restaurant_id=casa).first().photo_ref = \
            str(tmp_path / "la-que-ya-no-esta.jpg")

    destino = tmp_path / "todo.zip"
    with db.session_scope() as s:
        cuenta = exportar.paquete(s, casa, "es", destino)
    assert cuenta == {"fotos": 0, "faltan": 1, "bytes": 0}
    with zipfile.ZipFile(destino) as z:
        assert exportar.LIBRO in z.namelist()
        assert "1 fotos" in z.read(f"{exportar.FOTOS}/FALTAN.txt").decode()


def test_the_rows_say_which_photo_is_theirs(casas, tmp_path):
    """Sin eso, el que lo abre tiene las filas por un lado y las fotos por otro."""
    from thegrill.meat import exportar

    casa, _ = casas
    with db.session_scope() as s:
        pieza = s.query(Primal).filter_by(restaurant_id=casa).first()
        pieza.photo_ref = str(tmp_path / "loquesea.jpg")
        serial = pieza.serial
        assert exportar.nombre_foto(pieza) == f"{exportar.FOTOS}/{serial}.jpg"
    with db.session_scope() as s:
        libro = _abre(exportar.libro(s, casa, "es"))
    hoja = libro[libro.sheetnames[0]]
    cabeceras = [c.value for c in hoja[1]]
    assert "Foto de la etiqueta" in cabeceras, cabeceras


def test_the_house_that_stopped_paying_can_still_take_its_proof(casas, tmp_path):
    """Y por la puerta por la que se descarga, el día que ya no es cliente.

    Es el día que importa: el programa se cierra y lo que la ley le exige
    conservar sigue siendo suyo. Con las etiquetas, que son la prueba.
    """
    import zipfile

    from thegrill.meat import exportar

    casa, _ = casas
    carpeta = tmp_path / "subidas"
    carpeta.mkdir(exist_ok=True)
    with db.session_scope() as s:
        for pieza in s.query(Primal).filter_by(restaurant_id=casa):
            foto = carpeta / f"{pieza.serial}.jpg"
            foto.write_bytes(b"\xff\xd8\xff-etiqueta")
            pieza.photo_ref = str(foto)
    _bloquea(casa)

    cliente = entra("albano@marina.com")
    respuesta = cliente.get("/descargas/mis-datos.zip")
    assert respuesta.status_code == 200
    assert respuesta.headers["content-type"] == "application/zip"
    assert respuesta.headers["content-disposition"].endswith('.zip"')

    destino = tmp_path / "bajado.zip"
    destino.write_bytes(respuesta.content)
    with zipfile.ZipFile(destino) as z:
        assert exportar.LIBRO in z.namelist()
        assert any(n.startswith(f"{exportar.FOTOS}/") for n in z.namelist()), z.namelist()


def test_nobody_else_takes_the_house_away_in_a_zip_either(casas):
    """La puerta nueva tiene la misma cerradura que la de al lado."""
    casa, _ = casas
    de_la_plantilla = add_user("albano@marina.com", "cocinero@marina.com", "Luis")
    assert de_la_plantilla.get("/descargas/mis-datos.zip").status_code in (303, 403)
