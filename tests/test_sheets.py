"""Hojas en Excel para imprimir y rellenar a mano."""
import io
import xml.dom.minidom
import zipfile
from datetime import date

import pytest
from openpyxl import load_workbook

from thegrill import db
from thegrill.models import FieldType, RecordTemplate, TemplateField
from thegrill.web import auth, sheets
from thegrill.web.seed import seed_templates


@pytest.fixture
def ctx(tmp_path):
    db.init_engine(f"sqlite:///{tmp_path/'s.db'}")
    db.create_all()
    with db.session_scope() as s:
        rest, _ = auth.create_restaurant(s, "Bistró del Puerto", "ana@b.com", "Ana",
                                         "clave-larga-1", language="es")
        seed_templates(s, rest.id, "es")
        yield s, rest


def tpl(session, rest, code):
    return session.query(RecordTemplate).filter_by(restaurant_id=rest.id, code=code).one()


def sheet_of(payload: bytes, index: int = 0):
    return load_workbook(io.BytesIO(payload))[load_workbook(io.BytesIO(payload)).sheetnames[index]]


def test_the_file_is_a_valid_workbook(ctx):
    s, rest = ctx
    payload = sheets.workbook_for([tpl(s, rest, "merma")], rest, "es")
    z = zipfile.ZipFile(io.BytesIO(payload))
    assert z.testzip() is None
    for name in z.namelist():
        if name.endswith((".xml", ".rels")):
            xml.dom.minidom.parseString(z.read(name))    # falla si está mal formado


def test_columns_match_the_form_fields(ctx):
    s, rest = ctx
    fridge = tpl(s, rest, "temp_refrigeracion")
    ws = sheet_of(sheets.workbook_for([fridge], rest, "es"))
    headers = [c.value for c in ws[8]]
    assert headers[:3] == ["Nº", "Fecha", "Hora"]        # el papel necesita día y hora por fila
    assert len(headers) == 3 + len(fridge.fields)
    assert headers[3] == "Equipo *"                      # el asterisco marca lo obligatorio
    assert headers[4] == "Temperatura (°C) límites -2 a 5 *"
    assert headers[5] == "Acción correctiva si está fuera de rango"   # opcional, sin asterisco


def test_a_field_with_only_one_limit_says_so(ctx):
    s, rest = ctx
    t = RecordTemplate(restaurant_id=rest.id, code="x", name="X", category="other")
    t.fields.append(TemplateField(key="a", label="Solo mínimo", type=FieldType.NUMBER,
                                  unit="kg", min_value=2, required=False))
    t.fields.append(TemplateField(key="b", label="Solo máximo", type=FieldType.NUMBER,
                                  max_value=90, required=False))
    s.add(t)
    s.flush()
    headers = [c.value for c in sheet_of(sheets.workbook_for([t], rest, "es"))[8]]
    assert headers[3] == "Solo mínimo (kg) mínimo 2"
    assert headers[4] == "Solo máximo máximo 90"


def test_the_example_row_shows_the_expected_format(ctx):
    s, rest = ctx
    ws = sheet_of(sheets.workbook_for([tpl(s, rest, "recepcion")], rest, "es"))
    example = [c.value for c in ws[9]]
    assert example[0] == "EJEMPLO"
    assert example[1] == date.today().isoformat() and example[2] == "08:30"
    values = dict(zip([c.value for c in ws[8]][3:], example[3:]))
    assert values["Proveedor *"] == "—"                            # texto libre
    assert values["Cantidad (kg) *"] == "12"                       # número sin límites
    assert values["Fecha de caducidad *"] == date.today().isoformat()
    assert values["Mercancía conforme *"] == "Sí"                  # sí/no
    assert values["Temperatura de llegada (°C) límites -30 a 5"] == "-12.5"   # punto medio


def test_the_example_of_a_list_field_is_one_of_its_options(ctx):
    s, rest = ctx
    waste = tpl(s, rest, "merma")
    ws = sheet_of(sheets.workbook_for([waste], rest, "es"))
    motivo = next(f for f in waste.fields if f.key == "motivo")
    assert [c.value for c in ws[9]][5] in motivo.options.split("|")


def test_blank_rows_are_numbered_and_bordered(ctx):
    s, rest = ctx
    ws = sheet_of(sheets.workbook_for([tpl(s, rest, "limpieza")], rest, "es", blank_rows=5))
    assert [ws.cell(row=10 + i, column=1).value for i in range(5)] == [1, 2, 3, 4, 5]
    assert ws.cell(row=10, column=2).value is None
    assert ws.cell(row=10, column=2).border.left.style == "thin"
    assert "Firma del responsable" in ws.cell(row=16, column=1).value


def test_the_legend_tells_people_what_to_fill_and_what_to_ignore(ctx):
    s, rest = ctx
    ws = sheet_of(sheets.workbook_for([tpl(s, rest, "merma")], rest, "es"))
    assert ws["A1"].value == "Merma y desperdicio"
    assert rest.name in ws["A2"].value
    assert "no la uses" in ws["A6"].value and "obligatorio" in ws["A6"].value
    assert ws.cell(row=4, column=1).value == "Restaurante:"
    assert ws.cell(row=4, column=3).value == "Turno:"
    assert ws.cell(row=4, column=5).value == "Rellenado por:"


def test_it_is_set_up_to_print(ctx):
    s, rest = ctx
    ws = sheet_of(sheets.workbook_for([tpl(s, rest, "recepcion")], rest, "es"))
    assert str(ws.page_setup.paperSize) == str(ws.PAPERSIZE_A4)   # openpyxl lo relee como texto
    assert ws.page_setup.orientation == "landscape"       # siete columnas no caben en vertical
    assert ws.page_setup.fitToWidth == 1 and ws.page_setup.fitToHeight == 0
    assert ws.sheet_properties.pageSetUpPr.fitToPage is True
    assert ws.print_title_rows == "$1:$8"                 # el encabezado se repite en cada hoja
    assert ws.print_area.endswith("$J$31") if isinstance(ws.print_area, str) \
        else ws.print_area[0].endswith("$J$31")
    assert ws.sheet_view.showGridLines is False


def test_orientation_follows_how_wide_the_form_is(ctx):
    """Vertical mientras quepa, porque entran más filas por página."""
    s, rest = ctx
    narrow = RecordTemplate(restaurant_id=rest.id, code="n", name="Corto", category="other")
    narrow.fields.append(TemplateField(key="a", label="Lectura", type=FieldType.NUMBER))
    s.add(narrow)
    s.flush()
    assert sheet_of(sheets.workbook_for([narrow], rest, "es")).page_setup.orientation == "portrait"
    wide = sheet_of(sheets.workbook_for([tpl(s, rest, "recepcion")], rest, "es"))
    assert wide.page_setup.orientation == "landscape"


def test_one_tab_per_form_with_legal_names(ctx):
    s, rest = ctx
    everything = s.query(RecordTemplate).filter_by(restaurant_id=rest.id).all()
    names = load_workbook(io.BytesIO(sheets.workbook_for(everything, rest, "es"))).sheetnames
    assert len(names) == len(everything)
    assert "Producción   mise en place" in names          # la barra está prohibida en Excel
    assert all(len(n) <= 31 and not set(n) & set(":\\/?*[]") for n in names)


def test_repeated_or_very_long_names_do_not_collide(ctx):
    used = set()
    assert sheets.sheet_title("Control", used) == "Control"
    assert sheets.sheet_title("control", used) == "control 2"
    assert sheets.sheet_title("Control", used) == "Control 3"
    long_name = "Control de temperatura de las cámaras frigoríficas del sótano"
    first = sheets.sheet_title(long_name, used)
    second = sheets.sheet_title(long_name, used)
    assert len(first) == 31 and len(second) <= 31 and first != second


def test_an_empty_workbook_is_still_valid(ctx):
    s, rest = ctx
    names = load_workbook(io.BytesIO(sheets.workbook_for([], rest, "es"))).sheetnames
    assert names == ["—"]                                  # Excel rechaza un libro sin hojas


def test_the_sheet_speaks_the_readers_language(ctx):
    s, rest = ctx
    fridge = tpl(s, rest, "temp_refrigeracion")
    ws = sheet_of(sheets.workbook_for([fridge], rest, "de"))
    assert [c.value for c in ws[8]][:3] == ["Nº", "Datum", "Uhrzeit"]
    assert ws.cell(row=4, column=3).value == "Schicht:"
    assert [c.value for c in ws[9]][0] == "BEISPIEL"
    # las etiquetas de los campos son datos del restaurante, no se traducen al vuelo
    assert [c.value for c in ws[8]][3] == "Equipo *"


def test_arabic_sheets_read_right_to_left(ctx):
    s, rest = ctx
    payload = sheets.workbook_for([tpl(s, rest, "merma")], rest, "ar")
    assert sheet_of(payload).sheet_view.rightToLeft is True
    assert 'rightToLeft="1"' in zipfile.ZipFile(io.BytesIO(payload)).read(
        "xl/worksheets/sheet1.xml").decode()
    assert sheet_of(sheets.workbook_for([tpl(s, rest, "merma")], rest, "es")).sheet_view.rightToLeft is False


def test_filenames_are_safe(ctx):
    assert sheets.filename_for("temp_refrigeracion") == "temp_refrigeracion.xlsx"
    assert sheets.filename_for("../../etc/passwd") == ".._.._etc_passwd.xlsx"
    assert sheets.filename_for("مطعم النخيل") == "hoja.xlsx"     # sin letras latinas, nombre neutro
