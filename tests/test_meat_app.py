"""La edición solo carne: control de carnes para restaurantes y hoteles.

Lo que se comprueba es el recorrido entero de una pieza en esta app: llega en
un lote de recepción con su número y su coste, se despieza en cortes con serial
propio, sale a descongelar, se cuenta al cerrar el turno, se vende, se cuadra en
el inventario y se puede seguir su historia. Y que no hay ninguna puerta a la
cocina general: aquí solo hay carne.
"""
import os
import re
from datetime import date, timedelta

import pytest
from fastapi.testclient import TestClient

from thegrill import db
from thegrill.meat import app as meatapp
from thegrill.models import (Despiece, DespieceCut, Ingredient, IngredientItem,
                             IngredientLot,
                             IngredientMovement, MovementKind, Primal, PrimalStatus,
                             Recipe, Restaurant, Role, User)

HOY = date.today()


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("GRILL_INSECURE_COOKIE", "1")
    db.init_engine(f"sqlite:///{tmp_path/'carne.db'}")
    db.create_all()
    with TestClient(meatapp.app, follow_redirects=False, headers=SPANISH) as c:
        yield c


from tests.meat_helpers import add_user, csrf_from, join_code, login  # noqa: E402

# El navegador de estas pruebas habla español: los textos que se comprueban
# abajo son los españoles. Quien llega sin decir nada recibe inglés.
SPANISH = {"accept-language": "es"}


def signup(client, restaurant="Hotel Marina", email="albano@marina.com", name="Albano",
           language=""):
    """La casa la da de alta la plataforma; aquí entra su manager."""
    from tests.meat_helpers import signup as alta
    return alta(client, restaurant, email, name, language or "es")


def join(client, code=None, email="marta@marina.com", name="Marta"):
    """Ya no se entra con código: el manager crea la cuenta y se entra con ella."""
    add_user(client, email=email, name=name)
    return login(client, email, "clave-larga-2")


# ------------------------------------------------------ puertas y encuadre
def test_the_front_door_leads_to_today_not_to_the_kitchen(client):
    assert client.get("/").status_code == 200          # la portada pública
    signup(client)
    assert client.get("/").headers["location"] == "/hoy"
    assert client.get("/hoy").status_code == 200


def test_this_edition_says_what_it_is(client):
    signup(client)
    html = client.get("/hoy").text
    assert "Control de carnes" in html
    assert "Gestión de cocina" not in html


def test_there_is_no_door_to_anything_that_is_not_meat(client):
    """Ni recetas de cocina, ni plantillas HACCP, ni registros generales."""
    signup(client)
    for path in ("/recetas", "/manager/plantillas", "/manager/registros",
                 "/app", "/app/mis-registros", "/manager"):
        assert client.get(path).status_code == 404, path


def test_the_bar_only_shows_meat(client):
    signup(client)
    nav = client.get("/hoy").text
    for path in ("/recepcion", "/despiece", "/carne", "/descongelado", "/inventario",
                 "/merma", "/trazabilidad", "/cortes", "/carta", "/ingredientes"):
        assert f'href="{path}"' in nav, path
    assert 'href="/recetas"' not in nav          # las recetas de cocina no están aquí


def test_someone_who_already_has_an_account_can_come_back_in(client):
    """La puerta de casa: si esto se rompe, no entra nadie."""
    signup(client)
    client.cookies.clear()
    assert client.get("/hoy").headers["location"] == "/login"

    r = client.post("/login", data={"email": "albano@marina.com", "password": "clave-larga-1"})
    assert r.status_code == 303 and r.headers["location"] == "/hoy"
    assert client.get("/hoy").status_code == 200

    fuera = client.post("/logout")
    assert fuera.status_code == 303 and fuera.headers["location"] == "/login"
    client.cookies.clear()
    assert client.get("/hoy").headers["location"] == "/login"


def test_a_wrong_password_says_nothing_about_whether_the_email_exists(client):
    signup(client)
    client.cookies.clear()
    malo = client.post("/login", data={"email": "albano@marina.com", "password": "otra-clave"})
    nadie = client.post("/login", data={"email": "nadie@marina.com", "password": "otra-clave"})
    assert malo.status_code == nadie.status_code == 200
    assert "banner bad" in malo.text and "banner bad" in nadie.text
    assert client.get("/hoy").headers["location"] == "/login"


def test_an_employee_joins_with_the_code_and_lands_on_today(client):
    signup(client)
    client.cookies.clear()
    join(client, join_code())
    assert client.get("/hoy").status_code == 200


def test_health_says_which_edition_it_is(client):
    assert client.get("/healthz").json() == {"status": "ok", "edition": "meat"}


def test_anonymous_gets_sent_to_login(client):
    for path in ("/hoy", "/recepcion", "/despiece", "/carne", "/descongelado",
                 "/inventario", "/merma", "/trazabilidad", "/cortes", "/carta", "/ventas"):
        r = client.get(path)
        assert r.status_code == 303 and r.headers["location"] == "/login", path


# -------------------------------------------------------------- recepción
def test_a_delivery_books_each_piece_with_its_own_number_and_cost(client):
    signup(client)
    form = client.get("/recepcion")
    r = client.post("/recepcion", data={
        "csrf": csrf_from(form.text), "lot": "DXB20260910", "sku": "Striploin AUS",
        "grade": "MB9+", "origin": "AUS", "price:0": "32", "use_by": str(HOY + timedelta(days=40)),
        "serial:0": "8017", "g:0": "9400",
        "serial:1": "8018", "g:1": "10200", "price:1": "34"})
    # Guardar contesta con una redirección: recargar no da de alta otra vez.
    assert r.status_code == 303
    with db.session_scope() as s:
        piezas = s.query(Primal).order_by(Primal.serial).all()
        assert [p.serial for p in piezas] == ["8017", "8018"]
        assert all(p.lot == "DXB20260910" for p in piezas)      # el lote es común
        assert piezas[0].piece_cost_usd == round(9.4 * 32, 4)   # el coste, de cada una
        assert piezas[1].piece_cost_usd == round(10.2 * 34, 4)  # y puede ser distinto
        assert piezas[0].grade == "MB9+" and piezas[0].origin == "AUS"
        assert piezas[0].status == PrimalStatus.IN_STOCK


def test_two_pieces_can_never_share_a_number(client):
    signup(client)
    form = client.get("/recepcion")
    r = client.post("/recepcion", data={
        "csrf": csrf_from(form.text), "lot": "L1", "price:0": "30",
        "serial:0": "8017", "g:0": "9000",
        "serial:1": "8017", "g:1": "9000"})
    assert "8017" in r.text
    with db.session_scope() as s:
        assert s.query(Primal).count() == 0       # o entra el lote entero, o no entra nada


def test_a_piece_without_a_weight_is_refused(client):
    signup(client)
    form = client.get("/recepcion")
    r = client.post("/recepcion", data={"csrf": csrf_from(form.text), "lot": "L1",
                                        "serial:0": "8017", "g:0": ""})
    assert r.status_code == 200
    with db.session_scope() as s:
        assert s.query(Primal).count() == 0


def test_reception_needs_a_csrf_token(client):
    signup(client)
    assert client.post("/recepcion", data={"lot": "L1", "serial:0": "8017",
                                           "g:0": "9000"}).status_code == 403


# ---------------------------------------------------------------- cortes
def setup_cuts(client, consumption="RECIPE"):
    """Los cortes que va a dar el striploin, con su artículo de cámara.

    `consumption` dice de dónde sale su consumo: de la venta en el POS o del
    recuento de descongelado. Las dos cosas a la vez descontarían el doble.
    """
    form = client.get("/cortes")
    token = csrf_from(form.text)
    ids = {}
    for name in ("Striploin steak", "Tiras de striploin", "Recorte de vacuno"):
        client.post("/cortes/nuevo", data={"csrf": token, "name": name, "rotation": "FEFO",
                                           "consumption": consumption})
    with db.session_scope() as s:
        for cut in s.query(Ingredient).all():
            ids[cut.name] = cut.id
    for name, cut_id in ids.items():
        client.post(f"/cortes/{cut_id}/articulo",
                    data={"csrf": token, "name": f"{name} AUS", "supplier": "Aussie Meat"})
    with db.session_scope() as s:
        return {i.ingredient.name: i.id for i in s.query(IngredientItem).all()}


def test_a_cut_groups_articles_of_any_origin_in_one_queue(client):
    signup(client)
    setup_cuts(client)
    with db.session_scope() as s:
        cut = s.query(Ingredient).filter_by(name="Striploin steak").one()
        assert cut.category == "carne"
        assert [i.name for i in cut.items] == ["Striploin steak AUS"]
    assert "Striploin steak" in client.get("/cortes").text


def test_an_employee_cannot_invent_cuts(client):
    """El catálogo lo ve todo el equipo; tocarlo, solo dirección."""
    signup(client)
    client.cookies.clear()
    join(client, join_code())

    pagina = client.get("/cortes")
    assert pagina.status_code == 200                       # verlos sí
    assert 'action="/cortes/nuevo"' not in pagina.text     # pero sin formulario

    token = csrf_from(client.get("/merma").text)           # con un token válido suyo
    assert client.post("/cortes/nuevo",
                       data={"csrf": token, "name": "Solomillo"}).status_code == 403


# -------------------------------------------------------------- despiece
def deliver(client, serials=("8017",), kg=10.0, price=30.0):
    form = client.get("/recepcion")
    data = {"csrf": csrf_from(form.text), "lot": "DXB20260910", "sku": "Striploin AUS",
            "grade": "MB9+", "origin": "AUS",
            "use_by": str(HOY + timedelta(days=40))}
    # El precio del kilo es de cada pieza: no hay uno del camión que valga para
    # todas, porque dos bolsas del mismo camión no valen lo mismo.
    for i, serial in enumerate(serials):
        data[f"serial:{i}"] = serial
        data[f"kg:{i}"] = str(kg)
        data[f"price:{i}"] = str(price)
    # Guardar contesta con una redirección: el recado espera en la pantalla
    # de después, para que recargar no vuelva a dar de alta el camión.
    assert client.post("/recepcion", data=data).status_code == 303


def butcher(client, items, tg="TG-0001", before=10.0, waste="600"):
    form = client.get("/despiece")
    data = {"csrf": csrf_from(form.text), "tg": tg, "date": str(HOY),
            "staff": "Albano", "before_g": str(int(before * 1000)), "waste_g": waste,
            "primal": "8017",
            "cut:0": "Striploin steak", "item:0": items["Striploin steak"],
            "pieces:0": "20", "grams:0": "250", "index:0": "1,6",
            "cut:1": "Tiras", "item:1": items["Tiras de striploin"],
            "pieces:1": "8", "grams:1": "150", "index:1": "1",
            "cut:2": "Recorte", "item:2": items["Recorte de vacuno"],
            "pieces:2": "12", "grams:2": "200", "index:2": "0,45", "trim:2": "1"}
    return client.post("/despiece", data=data)


def test_butchering_turns_a_primal_into_cuts_with_their_own_serials(client):
    signup(client)
    items = setup_cuts(client)
    deliver(client)
    assert butcher(client, items).status_code == 303
    assert "TG-0001" in client.get("/despiece").text

    with db.session_scope() as s:
        lots = s.query(IngredientLot).order_by(IngredientLot.serial).all()
        assert [l.serial for l in lots] == ["8017-01", "8017-02", "8017-03"]
        assert all(l.parent_serial == "8017" for l in lots)     # de dónde salió cada uno
        assert all(l.lot_code == "TG-0001" for l in lots)
        assert all(l.grade == "MB9+" and l.origin == "AUS" for l in lots)
        # La pieza queda marcada como cortada, y solo por el despiece que lo dice.
        assert s.query(Primal).one().status == PrimalStatus.CUT
        assert s.query(Despiece).one().posted


def test_the_cost_of_the_primal_lands_on_the_cuts_and_the_waste_pays_nothing(client):
    signup(client)
    items = setup_cuts(client)
    deliver(client, kg=10.0, price=30.0)                   # 300 de coste
    butcher(client, items)
    with db.session_scope() as s:
        lots = s.query(IngredientLot).all()
        repartido = round(sum(l.qty * l.unit_cost for l in lots), 2)
        assert repartido == 300.0                          # no se pierde ni un céntimo
        filete = next(l for l in lots if l.serial == "8017-01")
        recorte = next(l for l in lots if l.serial == "8017-03")
        assert filete.unit_cost > recorte.unit_cost        # el filete vale más por kilo


def test_a_butchery_number_is_never_reused(client):
    signup(client)
    items = setup_cuts(client)
    deliver(client, serials=("8017", "8018"))
    butcher(client, items)
    r = butcher(client, items, tg="TG-0001")
    assert "TG-0001" in r.text
    with db.session_scope() as s:
        assert s.query(Despiece).count() == 1


def test_a_butchery_that_does_not_reconcile_leaves_nothing_behind(client):
    """Si el motor lo rechaza, no se queda un despiece a medias en la base."""
    signup(client)
    items = setup_cuts(client)
    deliver(client)
    form = client.get("/despiece")
    r = client.post("/despiece", data={
        "csrf": csrf_from(form.text), "tg": "TG-0009", "before_g": "10000",
        "primal": "9999",                                  # esa pieza no existe
        "cut:0": "Steak", "item:0": items["Striploin steak"],
        "pieces:0": "20", "grams:0": "250"})
    assert r.status_code == 200
    with db.session_scope() as s:
        assert s.query(Despiece).count() == 0
        assert s.query(IngredientLot).count() == 0


def test_a_rejected_butchery_never_loses_the_primal(client):
    """Lo que se tira es el papel, no la pieza: sigue entera y se puede volver a cortar."""
    signup(client)
    items = setup_cuts(client)
    deliver(client)

    form = client.get("/despiece")
    r = client.post("/despiece", data={
        "csrf": csrf_from(form.text), "tg": "TG-0007", "before_g": "10000", "primal": "8017",
        "cut:0": "Steak", "item:0": "", "pieces:0": "20", "grams:0": "250"})   # sin artículo
    assert r.status_code == 200

    with db.session_scope() as s:
        pieza = s.query(Primal).filter_by(serial="8017").one()
        assert pieza.status == PrimalStatus.IN_STOCK     # entera, en su sitio
        assert pieza.weight_kg == 10.0                   # con su peso
        assert pieza.piece_cost_usd == 300.0             # y con su coste
        assert s.query(Despiece).count() == 0            # solo se ha ido el papel

    # Y vuelve a estar en la lista, lista para despiezar otra vez.
    assert "8017" in client.get("/despiece").text
    assert butcher(client, items, tg="TG-0008").status_code == 303
    with db.session_scope() as s:
        assert s.query(Primal).filter_by(serial="8017").one().status == PrimalStatus.CUT
        assert s.query(IngredientLot).count() == 3


def test_only_a_posted_butchery_marks_a_primal_as_cut(client):
    """Nunca se da una pieza por cortada por inferencia: lo dice el despiece o nada."""
    signup(client)
    items = setup_cuts(client)
    deliver(client, serials=("8017", "8018", "8019"))
    butcher(client, items)                                # solo entra la 8017
    with db.session_scope() as s:
        estados = {p.serial: p.status for p in s.query(Primal)}
        assert estados["8017"] == PrimalStatus.CUT
        assert estados["8018"] == PrimalStatus.IN_STOCK
        assert estados["8019"] == PrimalStatus.IN_STOCK
    pantalla = client.get("/despiece").text
    assert "8018" in pantalla and "8019" in pantalla      # las otras dos siguen ahí


def test_a_piece_already_cut_cannot_be_butchered_twice(client):
    signup(client)
    items = setup_cuts(client)
    deliver(client)
    butcher(client, items)
    r = butcher(client, items, tg="TG-0050")
    assert r.status_code == 200 and "8017" in r.text
    with db.session_scope() as s:
        assert s.query(Despiece).count() == 1
        assert s.query(IngredientLot).count() == 3        # no se han duplicado los cortes


def test_a_butchery_whose_weight_does_not_add_up_is_posted_with_a_warning(client):
    """Un descuadre de masa avisa, pero no bloquea: la carne ya está cortada."""
    signup(client)
    items = setup_cuts(client)
    deliver(client)
    form = client.get("/despiece")
    r = client.post("/despiece", data={
        "csrf": csrf_from(form.text), "tg": "TG-0011", "before_g": "10000", "waste_g": "0",
        "primal": "8017",
        "cut:0": "Striploin steak", "item:0": items["Striploin steak"],
        "pieces:0": "10", "grams:0": "250"})              # 2,5 kg de 10: faltan 7,5
    assert r.status_code == 303
    assert "banner warn" in client.get("/despiece").text  # lo dice
    with db.session_scope() as s:
        assert s.query(Despiece).one().posted             # pero lo vuelca
        assert s.query(Primal).one().status == PrimalStatus.CUT


def test_a_butchery_can_never_borrow_another_kitchens_article(client):
    """Sin esto, un formulario manipulado cuelga un lote del corte del vecino."""
    signup(client)
    setup_cuts(client)
    deliver(client)
    client.cookies.clear()
    signup(client, restaurant="Otro hotel", email="bea@otro.com", name="Bea")
    ajeno = setup_cuts(client)["Striploin steak"]          # artículo de la otra casa
    client.cookies.clear()
    client.post("/login", data={"email": "albano@marina.com", "password": "clave-larga-1"})

    form = client.get("/despiece")
    r = client.post("/despiece", data={
        "csrf": csrf_from(form.text), "tg": "TG-0012", "before_g": "10000", "primal": "8017",
        "cut:0": "Steak", "item:0": ajeno, "pieces:0": "20", "grams:0": "250"})
    assert r.status_code == 200
    with db.session_scope() as s:
        assert s.query(IngredientLot).count() == 0
        assert s.query(Primal).filter_by(serial="8017").one().status == PrimalStatus.IN_STOCK


def test_a_cut_without_its_article_is_refused(client):
    signup(client)
    setup_cuts(client)
    deliver(client)
    form = client.get("/despiece")
    r = client.post("/despiece", data={
        "csrf": csrf_from(form.text), "tg": "TG-0002", "before_g": "10000", "primal": "8017",
        "cut:0": "Steak", "item:0": "", "pieces:0": "20", "grams:0": "250"})
    assert r.status_code == 200
    with db.session_scope() as s:
        assert s.query(Despiece).count() == 0


# ---------------------------------------------------------- descongelado
def test_the_shift_count_turns_into_real_consumption(client):
    signup(client)
    items = setup_cuts(client, consumption="COUNT")
    deliver(client)
    butcher(client, items)

    form = client.get("/descongelado")
    token = csrf_from(form.text)
    client.post("/descongelado/salida", data={"csrf": token, "serial": "8017-01",
                                              "pieces": "10", "total_g": "2500"})
    client.post("/descongelado/recuento", data={"csrf": token, "serial": "8017-01",
                                                "pieces": "4", "total_g": "1000"})
    estado = client.get("/descongelado").text
    assert "8017-01" in estado

    r = client.post("/descongelado/cierre", data={"csrf": token})
    assert r.status_code == 303
    with db.session_scope() as s:
        lot = s.query(IngredientLot).filter_by(serial="8017-01").one()
        assert lot.qty_remaining == pytest.approx(5.0 - 1.5, abs=0.001)   # salió 2,5, quedó 1,0


def test_what_went_out_but_was_not_counted_is_flagged_on_the_home_screen(client):
    signup(client)
    items = setup_cuts(client, consumption="COUNT")
    deliver(client)
    butcher(client, items)
    form = client.get("/descongelado")
    client.post("/descongelado/salida", data={"csrf": csrf_from(form.text), "serial": "8017-01",
                                              "pieces": "6", "total_g": "1500"})
    assert "recuento de cierre" in client.get("/hoy").text.lower()


# ---------------------------------------------------------------- carta
def test_a_meat_dish_is_a_cut_and_some_grams_tied_to_the_pos(client):
    signup(client)
    items = setup_cuts(client)
    deliver(client)
    butcher(client, items)
    with db.session_scope() as s:
        cut_id = s.query(Ingredient).filter_by(name="Striploin steak").one().id

    form = client.get("/carta")
    r = client.post("/carta/nuevo", data={
        "csrf": csrf_from(form.text), "name": "Entrecot a la brasa", "cut_id": cut_id,
        "grams": "330", "sale_price": "29,50", "vat_pct": "10",
        "pos_code": "1201", "pos_name": "ENTRECOT"})
    assert r.status_code == 303

    html = client.get("/carta").text
    assert "Entrecot a la brasa" in html and "1201" in html
    with db.session_scope() as s:
        dish = s.query(Recipe).one()
        assert len(dish.lines) == 1
        assert dish.lines[0].qty == 0.33            # los gramos, en kilos
        assert dish.lines[0].ingredient_id == cut_id


def test_selling_takes_it_off_the_chiller_and_says_which_dish_it_went_to(client):
    signup(client)
    items = setup_cuts(client)
    deliver(client)
    butcher(client, items)
    with db.session_scope() as s:
        cut_id = s.query(Ingredient).filter_by(name="Striploin steak").one().id
    form = client.get("/carta")
    client.post("/carta/nuevo", data={
        "csrf": csrf_from(form.text), "name": "Entrecot", "cut_id": cut_id, "grams": "330",
        "sale_price": "29,50", "vat_pct": "10", "pos_code": "1201", "pos_name": "ENTRECOT"})

    ventas = client.get("/ventas")
    assert "ENTRECOT" in ventas.text
    r = client.post("/ventas", data={"csrf": csrf_from(ventas.text),
                                     "business_date": str(HOY), "units:ENTRECOT": "6"})
    assert r.status_code == 303
    with db.session_scope() as s:
        lot = s.query(IngredientLot).filter_by(serial="8017-01").one()
        assert lot.qty_remaining == pytest.approx(5.0 - 6 * 0.33, abs=0.001)
        # Cada salida deja escrito a qué plato fue: sin eso no hay food cost por pieza.
        salidas = s.query(IngredientMovement).filter_by(kind=MovementKind.SALE).all()
        assert salidas and all(m.source_ref == "ENTRECOT" for m in salidas)


# ------------------------------------------------------------ trazabilidad
def test_the_history_of_a_piece_reads_from_delivery_to_plate(client):
    signup(client)
    items = setup_cuts(client)
    deliver(client)
    butcher(client, items)
    html = client.get("/trazabilidad?serial=8017").text
    assert "TG-0001" in html
    assert "8017-01" in html


# --------------------------------------------------------------- aislamiento
def test_one_kitchen_never_sees_another_hotels_meat(client):
    signup(client)
    items = setup_cuts(client)
    deliver(client)
    butcher(client, items)

    client.cookies.clear()
    signup(client, restaurant="Otro hotel", email="bea@otro.com", name="Bea")
    assert "8017" not in client.get("/recepcion").text
    assert "8017-01" not in client.get("/carne").text
    assert client.get("/trazabilidad?serial=8017").text.count("8017-01") == 0


# ---------------------------------------------------------------- idiomas
def test_the_meat_screens_speak_the_six_languages(client):
    signup(client, language="de")
    html = client.get("/hoy").text
    assert "Fleischkontrolle" in html
    assert "Zerlegung" in client.get("/despiece").text
    assert 'dir="ltr"' in html


def test_arabic_turns_the_whole_thing_right_to_left(client):
    signup(client, language="ar")
    html = client.get("/hoy").text
    assert 'dir="rtl"' in html
    assert "مراقبة اللحوم" in html


def test_no_meat_screen_leaks_spanish_when_the_language_is_english(client):
    signup(client, restaurant="Marina Hotel", email="ann@marina.com", name="Ann",
           language="en")
    spanish = ["Recepción", "Despiece", "Cámara", "Descongelado", "Cortes",
               "Configuración", "Salir", "Avisos", "Inventario", "Merma",
               "Trazabilidad", "Descargas"]
    for path in ("/hoy", "/recepcion", "/despiece", "/carne", "/descongelado",
                 "/inventario", "/merma", "/trazabilidad", "/cortes", "/carta",
                 "/ventas", "/descargas", "/configuracion", "/manager/equipo"):
        html = client.get(path).text
        leaked = [w for w in spanish if w in html]
        assert not leaked, f"{path} deja en español: {leaked}"


# --------------------------------------------------------------- descargas
def test_the_printable_sheets_are_meat_sheets(client):
    import io

    from openpyxl import load_workbook
    signup(client)
    html = client.get("/descargas").text
    for code in ("recepcion", "despiece", "descongelado", "inventario", "merma"):
        assert f"/descargas/{code}.xlsx" in html
        payload = client.get(f"/descargas/{code}.xlsx").content
        ws = load_workbook(io.BytesIO(payload)).active
        assert ws.page_setup.fitToWidth == 1
        assert ws.print_title_rows.replace("$", "") == "1:8"


def test_everything_in_one_workbook_has_one_tab_per_sheet(client):
    import io

    from openpyxl import load_workbook
    signup(client)
    from thegrill.meat import sheets_meat
    wb = load_workbook(io.BytesIO(client.get("/descargas/todo.xlsx").content))
    assert len(wb.sheetnames) == len(sheets_meat.SHEETS) == 6


def test_an_unknown_sheet_is_not_a_download(client):
    signup(client)
    assert client.get("/descargas/recetas.xlsx").status_code == 404


# ------------------------------------------------- los seis idiomas, enteros
def test_no_language_is_left_with_untranslated_meat_keys():
    """Añadir un idioma es copiar el español y traducirlo. Nada a medias."""
    from thegrill.meat import i18n_meat
    reference = set(i18n_meat.ES)
    for code, table in i18n_meat.CATALOGUES.items():
        assert set(table) == reference, f"{code} no cuadra con el español"
        assert all(str(v).strip() for v in table.values()), f"{code} deja textos vacíos"


def test_a_placeholder_in_one_language_is_a_placeholder_in_all_of_them():
    from thegrill.meat import i18n_meat
    for key, spanish in i18n_meat.ES.items():
        marks = set(re.findall(r"\{(\w+)\}", spanish))
        for code, table in i18n_meat.CATALOGUES.items():
            assert set(re.findall(r"\{(\w+)\}", table[key])) == marks, f"{code} · {key}"


def test_the_meat_texts_do_not_step_on_the_kitchen_ones():
    """Los textos de carne se añaden a los de cocina, no los reescriben."""
    import pathlib

    from thegrill.meat import i18n_meat
    # Las claves de la cocina se leen del fichero: en memoria ya están mezcladas.
    fuente = pathlib.Path("thegrill/web/i18n.py").read_text(encoding="utf-8")
    trozo = fuente[fuente.index("ES = {"):fuente.index("\n}", fuente.index("ES = {"))]
    cocina = set(re.findall(r'^\s{4}"([^"]+)":', trozo, re.M))
    pisadas = cocina & set(i18n_meat.ES)
    assert not pisadas, f"la edición de carne reescribe textos de cocina: {pisadas}"


def test_the_front_page_says_which_program_this_is_before_logging_in(client):
    """Quien abre el enlace tiene que saber dónde entra, y en su idioma."""
    for path in ("/login", "/", "/precios", "/solicitar"):
        html = client.get(path).text
        assert "Control de carnes" in html, path
        assert "Gestión de cocina" not in html, path
    ingles = client.get("/login", headers={"accept-language": "en"}).text
    assert "Meat control" in ingles and "Kitchen management" not in ingles


def test_there_is_a_printable_sheet_for_the_aging_fridge(client):
    """Se pesa de pie al lado de la nevera: esa hoja también se imprime."""
    import io

    from openpyxl import load_workbook
    from tests import meat_helpers as helpers

    helpers.signup(client)
    pantalla = client.get("/descargas")
    assert pantalla.status_code == 200
    assert "Maduración y congelador" in pantalla.text
    assert "/descargas/maduracion.xlsx" in pantalla.text

    hoja = client.get("/descargas/maduracion.xlsx")
    assert hoja.status_code == 200
    libro = load_workbook(io.BytesIO(hoja.content))
    cabecera = [c.value for c in libro.active[8]]
    assert "Peso de hoy (kg)" in cabecera and "Limpieza" in cabecera
    assert "Lo que se tira" in cabecera

    todas = client.get("/descargas/todo.xlsx")
    libro = load_workbook(io.BytesIO(todas.content))
    assert any("aduraci" in nombre for nombre in libro.sheetnames)


# --------------------------------------------- guiones que el navegador acepta
# Todas las pantallas de dentro. Si una trae un guion sin el número de la
# respuesta, el navegador lo tira sin decir nada: el botón no hace nada y no
# hay error que mirar. Pasó con el de activar los avisos.
CON_GUION = ["/hoy", "/carne", "/maduracion", "/descongelado", "/inventario",
             "/recepcion", "/despiece", "/merma", "/traslados", "/ventas",
             "/cortes", "/carta", "/notificaciones", "/configuracion",
             "/trazabilidad", "/parte", "/sedes", "/descargas", "/fallo"]


def test_no_screen_carries_a_script_the_browser_will_refuse(client):
    """Un guion sin el número de la respuesta no se ejecuta, y nadie se entera.

    Es el fallo más callado que hay: la página se ve entera, el botón está ahí
    y al pulsarlo no pasa nada. No hay error en pantalla ni en el servidor —el
    navegador lo bloquea por su cuenta—, así que solo se descubre probándolo a
    mano. Por eso se mira aquí pantalla por pantalla.
    """
    signup(client)
    sin_marca = []
    for ruta in CON_GUION:
        respuesta = client.get(ruta)
        assert respuesta.status_code == 200, f"{ruta}: {respuesta.status_code}"
        suyo = re.search(r"'nonce-([^']+)'",
                         respuesta.headers["content-security-policy"]).group(1)
        for etiqueta in re.findall(r"<script[^>]*>", respuesta.text):
            if "src=" in etiqueta:
                continue          # los de fichero valen por venir de casa
            if f'nonce="{suyo}"' not in etiqueta:
                sin_marca.append(f"{ruta}: {etiqueta[:60]}")
    assert not sin_marca, "guiones que el navegador va a tirar: " + "; ".join(sin_marca)


def test_the_button_for_browser_alerts_is_wired_up(client):
    """El botón de activar los avisos, con su guion y su marca."""
    signup(client)
    pantalla = client.get("/notificaciones")
    marca = re.search(r"'nonce-([^']+)'",
                      pantalla.headers["content-security-policy"]).group(1)
    assert 'id="askperm"' in pantalla.text
    assert f'<script nonce="{marca}">' in pantalla.text
    assert "Notification.requestPermission" in pantalla.text


def test_the_screens_tell_the_browser_which_theme_they_are_in(client):
    """Sin decirlo, el navegador pinta sus barras en blanco sobre lo oscuro.

    No es un capricho: la barra de desplazamiento, los desplegables y el
    calendario los pinta el navegador, no nosotros, y si no sabe en qué tema
    va los saca en claro. En una pantalla oscura eso es una raya de tiza al
    lado de la carne.
    """
    signup(client)
    for ruta in ("/hoy", "/carne", "/inventario", "/recepcion"):
        html = client.get(ruta).text
        # En la propia página, no en la hoja de estilo: la hoja es un fichero
        # aparte y llega unas décimas después. En esas décimas el navegador
        # saca las barras y los desplegables en claro, y luego cambian.
        assert '<meta name="color-scheme" content="light dark">' in html, ruta


# ------------------------------------------- dos momentos del día, dos pantallas
def test_taking_out_to_thaw_and_counting_at_closing_are_two_screens(client):
    """Los dos formularios son iguales: juntos, uno se escribe en el otro.

    Sacar a descongelar es de media mañana, con la cámara abierta; el recuento
    es de madrugada, al cerrar. Estaban a un palmo el uno del otro y con los
    mismos campos —serial, piezas, kilos, nota—: la salida acababa en la
    casilla del recuento y el turno salía descuadrado sin que nadie lo viera.
    """
    signup(client)
    salida = client.get("/descongelado").text
    assert '/descongelado/salida' in salida
    assert '/descongelado/recuento"' not in salida.split('class="tabsrow"')[1].split("</div>")[1]
    assert "/descongelado/cierre" not in salida       # no se cierra desde aquí

    recuento = client.get("/descongelado/recuento")
    assert recuento.status_code == 200
    assert '/descongelado/recuento' in recuento.text
    assert "/descongelado/cierre" in recuento.text    # el cierre va con el recuento
    assert '/descongelado/salida' not in recuento.text

    # Y se pasa de una a otra sin tener que buscarla.
    for pantalla in (salida, recuento.text):
        assert 'class="tabsrow"' in pantalla
        assert 'href="/descongelado"' in pantalla
        assert 'href="/descongelado/recuento"' in pantalla


def grupos_del_menu(html: str) -> dict:
    """Los grupos del menú lateral y si vienen abiertos."""
    return {m.group(1): bool(m.group(2))
            for m in re.finditer(r'<details class="grp" data-grp="([^"]+)"\s*(open)?>', html)}


def test_the_menu_groups_fold_up(client):
    """El menú entero no cabía en una ventana baja: ahora se pliega por grupos.

    Y se pliega solo el que no hace falta: el grupo de la pantalla en la que
    estás viene abierto siempre, porque plegarlo encima de lo que se está
    mirando es peor que no poder plegar nada.
    """
    signup(client)
    en_hoy = grupos_del_menu(client.get("/hoy").text)
    assert len(en_hoy) >= 4, en_hoy
    assert list(en_hoy.values())[0] is True            # el del día, abierto
    assert sum(1 for abierto in en_hoy.values() if not abierto) >= 2

    # Y en una pantalla de números se abre el suyo, no el de antes.
    en_inventario = grupos_del_menu(client.get("/inventario").text)
    assert en_inventario["Los números"] is True, en_inventario
    assert en_inventario["Catálogo"] is False, en_inventario
    assert en_inventario["La carne"] is False, en_inventario


# ------------------------------------------------- lo que se manda dos veces
def test_a_submission_that_arrives_twice_is_applied_once(client):
    """El teléfono reintenta cuando no sabe si el primero entró. Y a veces sí.

    Sin esto, el reintento da de alta el camión otra vez: las mismas piezas,
    duplicadas, con sus kilos y su dinero. El número del envío lo pone el
    teléfono antes del primer intento y es el mismo en todos los reintentos,
    así que el servidor reconoce al segundo y contesta que ya está hecho sin
    tocar nada.
    """
    signup(client)
    form = client.get("/recepcion")
    datos = {"csrf": csrf_from(form.text), "lot": "L-REPE", "sku": "Striploin AUS",
             "price:0": "30", "serial:0": "9001", "g:0": "9400",
             "serial:1": "9002", "g:1": "10200", "envio": "mismo-numero-de-envio"}

    primera = client.post("/recepcion", data=datos)
    segunda = client.post("/recepcion", data=datos)
    assert primera.status_code == 303
    assert segunda.status_code == 303          # «ya está hecho», sin escribir

    with db.session_scope() as s:
        piezas = [p.serial for p in s.query(Primal).order_by(Primal.serial)]
        assert piezas == ["9001", "9002"], piezas


def test_two_different_submissions_are_both_applied(client):
    """Y dos envíos distintos son dos: la protección no puede comerse trabajo."""
    signup(client)
    for numero, serial in (("uno", "9001"), ("dos", "9002")):
        form = client.get("/recepcion")
        client.post("/recepcion", data={
            "csrf": csrf_from(form.text), "lot": "L-DOS", "sku": "Striploin AUS",
            "price:0": "30", "serial:0": serial, "g:0": "9400", "envio": numero})
    with db.session_scope() as s:
        assert [p.serial for p in s.query(Primal).order_by(Primal.serial)] == ["9001", "9002"]


# ============================== lo que acaba de pasar, para el de al lado
class TestNovedades:
    """Dos personas trabajando la misma carne desde pantallas distintas.

    El del muelle da de alta seis lomos mientras el de la mesa despieza, y el
    que está contando no se entera de ninguna de las dos cosas hasta que va a
    la cámara. En FEFO eso se paga: se saca la pieza vieja porque nadie sabía
    que había entrado una nueva, o se cierra un inventario sin lo que entró
    hace diez minutos. Así que lo que pasa sale arriba, en la pantalla de quien
    esté trabajando, con su X para quitarlo cuando se ha leído.
    """

    def novedades(self, cliente, desde=0):
        r = cliente.get(f"/api/novedades?desde={desde}")
        assert r.status_code == 200, r.text[:200]
        return r.json()

    def test_a_delivery_tells_the_rest_of_the_house(self, client):
        signup(client)
        marta = add_user(client, email="marta@marina.com", name="Marta", role=Role.BUTCHER)
        deliver(marta, serials=("8017", "8018"), kg=9.4)

        avisa = self.novedades(client)
        assert len(avisa["items"]) == 1
        texto = avisa["items"][0]["texto"]
        assert "2 ×" in texto and "Striploin AUS" in texto      # cuántas y de qué
        assert "18,8 kg" in texto                               # los kilos que entran
        assert "DXB20260910" in texto                           # con qué lote: FEFO
        assert "Marta" in texto                                 # y quién

    def test_nobody_is_told_what_they_just_did_themselves(self, client):
        """Quien acaba de recibir ya sabe que ha recibido. El aviso es para los otros."""
        signup(client)
        marta = add_user(client, email="marta@marina.com", name="Marta", role=Role.BUTCHER)
        deliver(marta, serials=("8017",))

        assert self.novedades(marta)["items"] == []
        # Y aun así el número avanza: si no, el teléfono volvería a preguntar
        # por lo suyo cada treinta segundos y nunca pasaría de ahí.
        assert self.novedades(marta)["ultimo"] > 0

    def test_butchery_tells_the_rest_of_the_house(self, client):
        signup(client)
        items = setup_cuts(client)
        deliver(client)
        paco = add_user(client, email="paco@marina.com", name="Paco", role=Role.BUTCHER)
        butcher(paco, items)

        items_avisados = self.novedades(client)["items"]
        cortes = [x for x in items_avisados if x["kind"] == "DESPIECE"]
        assert len(cortes) == 1
        texto = cortes[0]["texto"]
        assert "TG-0001" in texto and "Striploin AUS" in texto
        assert "kg" in texto and "Paco" in texto

    def test_the_notice_is_read_in_the_language_of_whoever_reads_it(self, client):
        """El del muelle escribe en español y el jefe de cocina lo lee en inglés."""
        signup(client)
        marta = add_user(client, email="marta@marina.com", name="Marta", role=Role.BUTCHER)
        deliver(marta, serials=("8017",))

        assert "Han entrado" in self.novedades(client)["items"][0]["texto"]

        # El mismo hecho, leído por alguien que tiene la cuenta en inglés: se
        # guarda lo que pasó, no la frase, y la frase se arma al leerla.
        with db.session_scope() as s:
            s.query(User).filter_by(email="albano@marina.com").one().language = "en"
        assert "Just in" in self.novedades(client)["items"][0]["texto"]

    def test_what_is_already_read_does_not_come_back(self, client):
        signup(client)
        marta = add_user(client, email="marta@marina.com", name="Marta", role=Role.BUTCHER)
        deliver(marta, serials=("8017",))

        primera = self.novedades(client)
        leido = primera["items"][0]["id"]
        assert self.novedades(client, desde=leido)["items"] == []

        # Y lo que pase después sí vuelve a salir: cerrar un aviso no apaga el
        # siguiente.
        deliver(marta, serials=("8019",))
        siguiente = self.novedades(client, desde=leido)["items"]
        assert len(siguiente) == 1 and siguiente[0]["id"] > leido

    def test_yesterdays_news_is_not_news(self, client):
        """Quien entra por la mañana no quiere el turno de noche encima del título."""
        from datetime import datetime

        from thegrill.models import Novedad
        signup(client)
        marta = add_user(client, email="marta@marina.com", name="Marta", role=Role.BUTCHER)
        deliver(marta, serials=("8017",))
        with db.session_scope() as s:
            fila = s.query(Novedad).one()
            fila.created_at = datetime.utcnow() - timedelta(hours=20)

        assert self.novedades(client)["items"] == []

    def test_what_happens_at_another_site_does_not_interrupt_here(self, client):
        """Lo que entra en el obrador no le hace falta al local de la playa."""
        from thegrill.models import Site, SiteKind
        signup(client)
        client.post("/sedes/nueva", data={"name": "Playa", "kind": "OUTLET",
                                          "csrf": csrf_from(client.get("/sedes").text)})
        with db.session_scope() as s:
            playa = s.query(Site).filter_by(kind=SiteKind.OUTLET).one().id

        marta = add_user(client, email="marta@marina.com", name="Marta", role=Role.BUTCHER)
        deliver(marta, serials=("8017",))          # entra en la sede principal

        eva = add_user(client, email="eva@marina.com", name="Eva")
        with db.session_scope() as s:
            s.query(User).filter_by(email="eva@marina.com").one().site_id = playa

        assert self.novedades(eva)["items"] == []          # a Eva, en la playa, no
        assert len(self.novedades(client)["items"]) == 1   # al manager, que no tiene sede, sí

    def test_the_news_door_is_closed_from_outside(self, client):
        assert client.get("/api/novedades").headers["location"] == "/login"

    def test_old_news_is_forgotten(self, client):
        from datetime import datetime

        from thegrill.meat import novedades as mod
        from thegrill.models import Novedad
        signup(client)
        marta = add_user(client, email="marta@marina.com", name="Marta", role=Role.BUTCHER)
        deliver(marta, serials=("8017",))
        with db.session_scope() as s:
            s.query(Novedad).one().created_at = datetime.utcnow() - timedelta(days=9)
        with db.session_scope() as s:
            assert mod.olvidar_viejas(s) == 1
            assert s.query(Novedad).count() == 0

    def test_every_work_screen_can_show_them(self, client):
        """El aviso no sirve si solo sale en una pantalla: sale en todas."""
        from conftest import con_lo_de_fuera
        signup(client)
        for ruta in ("/hoy", "/recepcion", "/despiece", "/carne", "/inventario", "/merma"):
            html = client.get(ruta).text
            assert 'id="pilavisos"' in html, ruta
            assert "/api/novedades" in con_lo_de_fuera(client, html), ruta

    def test_a_critical_alert_can_tell_which_one_is_new(self, client):
        """El contador de avisos manda el número de cada uno.

        Sin él, la pantalla no sabe cuál acaba de llegar y la alerta crítica que
        debía saltar al teléfono no saltaba nunca.
        """
        signup(client)
        datos = client.get("/api/notificaciones").json()
        assert "items" in datos
        for aviso in datos["items"]:
            assert isinstance(aviso.get("id"), int)


# ================== la etiqueta del proveedor y el precio de dirección
class TestEtiquetaYPrecio:
    """De dónde viene cada pieza, y quién dice lo que vale.

    En el muelle se apunta lo que llega —qué es, cuánto pesa, de qué calidad y
    de dónde viene, con la etiqueta del proveedor delante— y se le hace una
    foto. El precio no lo sabe quien descarga, ni tiene por qué: el dinero es
    de dirección y además llega después, en la factura. Así que la pieza entra
    sin precio, se queda esperando, y no se puede despiezar hasta que alguien
    la activa.
    """

    def recibir(self, client, **extra):
        form = client.get("/recepcion")
        data = {"csrf": csrf_from(form.text), "lot": "L-ETQ", "sku": "Ribeye AUS",
                "grade": "MB7", "origin": "AUS",
                "producer_plant": "Teys Biloela", "est_code": "ES 10.00123/L",
                "breed": "Angus", "slaughter_date": str(HOY - timedelta(days=21)),
                "pack_date": str(HOY - timedelta(days=18)),
                "label_product": "CUBE ROLL GF", "halal": "1",
                "use_by": str(HOY + timedelta(days=40)),
                "serial:0": "8017", "g:0": "9400", "slot:0": "L-88213",
                "serial:1": "8018", "g:1": "9800", "slot:1": "L-88214",
                "grade:1": "MB9+", "slaughter:1": str(HOY - timedelta(days=20))}
        data.update(extra)
        return client.post("/recepcion", data=data)

    def test_each_piece_keeps_the_label_it_came_with(self, client):
        signup(client)
        assert self.recibir(client).status_code == 303
        with db.session_scope() as s:
            piezas = {p.serial: p for p in s.query(Primal)}
            uno, dos = piezas["8017"], piezas["8018"]
            # Lo que es igual para el camión se escribe una vez y se copia.
            assert uno.producer_plant == dos.producer_plant == "Teys Biloela"
            assert uno.est_code == "ES 10.00123/L" and uno.breed == "Angus"
            assert uno.label_product == "CUBE ROLL GF" and uno.halal is True
            assert uno.pack_date == HOY - timedelta(days=18)
            # Y lo que cambia de una bolsa a otra, va en su línea.
            assert uno.supplier_lot == "L-88213" and dos.supplier_lot == "L-88214"
            assert uno.grade == "MB7" and dos.grade == "MB9+"
            assert uno.slaughter_date == HOY - timedelta(days=21)
            assert dos.slaughter_date == HOY - timedelta(days=20)

    def test_a_label_date_that_could_not_have_happened_is_refused(self, client):
        """Un dedo en el teclado es peor que no tener el dato: se guarda y miente."""
        signup(client)
        r = self.recibir(client, slaughter_date=str(HOY + timedelta(days=3)))
        assert r.status_code == 200 and "8017" in r.text
        with db.session_scope() as s:
            assert s.query(Primal).count() == 0        # o entra todo, o no entra nada

    def test_meat_cannot_be_packed_before_it_is_slaughtered(self, client):
        signup(client)
        r = self.recibir(client, pack_date=str(HOY - timedelta(days=30)))
        assert r.status_code == 200
        with db.session_scope() as s:
            assert s.query(Primal).count() == 0

    def test_the_label_is_what_the_traceability_screen_opens_with(self, client):
        signup(client)
        self.recibir(client)
        ficha = client.get("/trazabilidad?serial=8017").text
        for dato in ("Teys Biloela", "ES 10.00123/L", "Angus", "L-88213", "CUBE ROLL GF"):
            assert dato in ficha, dato

    def test_a_recall_call_finds_the_pieces_by_what_the_caller_knows(self, client):
        """Quien llama para retirar algo no sabe nuestro número: sabe el suyo."""
        signup(client)
        self.recibir(client)
        for termino in ("L-88213", "Teys", "10.00123"):
            encontrado = client.get(f"/trazabilidad?serial={termino}").text
            assert "8017" in encontrado, termino

    # ------------------------------------------------ el precio y la activación
    def test_the_dock_does_not_set_prices(self, client):
        """El carnicero recibe; el dinero no es suyo y no se le enseña."""
        signup(client)
        paco = add_user(client, email="paco@marina.com", name="Paco", role=Role.BUTCHER)
        pantalla = paco.get("/recepcion").text
        assert 'name="price:0"' not in pantalla      # ni el de la pieza
        assert 'name="freight_kg"' not in pantalla   # ni lo que costó traerla
        # Y aunque lo escriba a mano en el formulario, no entra.
        self.recibir(paco, **{"price:0": "32"})
        with db.session_scope() as s:
            assert all(p.landed_usd_per_kg is None for p in s.query(Primal))

    def test_a_piece_without_a_price_cannot_be_butchered(self, client):
        """Repartir cero entre los cortes es perder el rastro del dinero."""
        signup(client)
        items = setup_cuts(client)
        paco = add_user(client, email="paco@marina.com", name="Paco", role=Role.BUTCHER)
        self.recibir(paco)

        # No sale en la lista de la pantalla…
        assert "8017" not in paco.get("/despiece").text
        # …y tampoco entra escribiéndola a mano.
        r = butcher(paco, items)
        assert "8017" in r.text
        with db.session_scope() as s:
            assert s.query(Despiece).count() == 0

    def test_management_is_told_there_are_prices_waiting(self, client):
        from thegrill.models import Notification

        signup(client)
        paco = add_user(client, email="paco@marina.com", name="Paco", role=Role.BUTCHER)
        self.recibir(paco)
        with db.session_scope() as s:
            avisos = s.query(Notification).all()
            assert len(avisos) == 1
            assert "2" in avisos[0].title
            assert "Paco" in avisos[0].body and "L-ETQ" in avisos[0].body
        # Y le aparece en el contador de la cabecera, que es donde mira.
        assert client.get("/api/notificaciones").json()["unread"] == 1

    def test_management_activates_the_pieces_and_then_they_can_be_cut(self, client):
        signup(client)
        items = setup_cuts(client)
        paco = add_user(client, email="paco@marina.com", name="Paco", role=Role.BUTCHER)
        self.recibir(paco)

        pantalla = client.get("/recepcion/precios")
        assert pantalla.status_code == 200
        assert "8017" in pantalla.text and "Teys Biloela" in pantalla.text
        puesto = client.post("/recepcion/precios", data={
            "csrf": csrf_from(pantalla.text), "serial": ["8017", "8018"],
            "all_price": "32", "price:8018": "34"})
        assert puesto.status_code == 200

        with db.session_scope() as s:
            piezas = {p.serial: p for p in s.query(Primal)}
            assert piezas["8017"].landed_usd_per_kg == 32       # el de todas
            assert piezas["8018"].landed_usd_per_kg == 34       # salvo el suyo
            assert piezas["8017"].piece_cost_usd == round(9.4 * 32, 4)
            assert piezas["8018"].piece_cost_usd == round(9.8 * 34, 4)

        # Y ahora sí se despieza.
        assert "8017" in paco.get("/despiece").text
        assert butcher(paco, items).status_code == 303
        with db.session_scope() as s:
            assert s.query(Despiece).count() == 1

    def test_a_price_of_zero_is_not_a_price(self, client):
        signup(client)
        paco = add_user(client, email="paco@marina.com", name="Paco", role=Role.BUTCHER)
        self.recibir(paco)
        pantalla = client.get("/recepcion/precios")
        r = client.post("/recepcion/precios", data={
            "csrf": csrf_from(pantalla.text), "serial": ["8017"], "price:8017": "0"})
        assert r.status_code == 200
        with db.session_scope() as s:
            assert s.query(Primal).filter_by(serial="8017").one().landed_usd_per_kg is None

    def test_only_management_touches_prices(self, client):
        signup(client)
        paco = add_user(client, email="paco@marina.com", name="Paco", role=Role.BUTCHER)
        assert paco.get("/recepcion/precios").status_code == 403
        assert paco.post("/recepcion/precios",
                         data={"csrf": "x", "serial": ["8017"]}).status_code == 403

    def test_the_vat_paid_on_the_purchase_is_never_a_cost(self, client):
        """El IVA de la compra se recupera: en el kilo no entra.

        Un lomo a 40 con un 10 % de IVA cuesta 44 de caja, pero **cuesta 40**.
        Esos 4 se descuentan del IVA que se cobra al vender: no son dinero que
        la casa pierda. Metidos en el kilo subirían el food cost de todos los
        platos que llevan ese lomo un diez por ciento, y a partir de ahí toda
        la carta estaría mal puesta.

        Se guarda igual, porque es lo que se declara.
        """
        signup(client)
        self.recibir(client, **{"price:0": "40", "price:1": "40",
                                "freight_kg": "2", "duty_kg": "1", "vat_pct": "10"})
        with db.session_scope() as s:
            una = s.query(Primal).filter_by(serial="8017").one()
            # El kilo es el de siempre: 40 + 2 de flete + 1 de aduana.
            assert una.landed_usd_per_kg == 43.0
            assert una.purchase_vat_pct == 10.0
            assert una.piece_cost_usd == round(9.4 * 43.0, 2)

    def test_the_report_adds_up_the_vat_that_can_be_deducted(self, client):
        """Y lo suma para la declaración: base, IVA y lo pagado de verdad."""
        from thegrill.meat import service as meat_service
        from thegrill.web import impuestos

        signup(client)
        self.recibir(client, **{"price:0": "40", "price:1": "40", "vat_pct": "10"})
        with db.session_scope() as s:
            piezas = s.query(Primal).all()
            iva = impuestos.soportado_de(piezas)
            assert iva.hay and iva.piezas == 2
            base = round(sum(p.piece_cost_usd for p in piezas), 2)
            assert iva.base == base
            assert iva.iva == round(base * 0.10, 2)
            assert iva.total == round(base * 1.10, 2)

            # Y sale en el parte del día que las recibió.
            parte = meat_service.daily_report(s, piezas[0].restaurant_id,
                                              on=piezas[0].received_date)
            assert len(parte.received) == 2
        assert "descontar" in client.get("/parte").text.lower()

    def test_a_purchase_without_a_vat_rate_is_not_guessed(self, client):
        """Sin apuntarlo no se supone ninguno: no es el mismo para el vino."""
        from thegrill.web import impuestos

        signup(client)
        self.recibir(client, **{"price:0": "40", "price:1": "40"})
        with db.session_scope() as s:
            piezas = s.query(Primal).all()
            assert all(p.purchase_vat_pct is None for p in piezas)
            assert not impuestos.soportado_de(piezas).hay

    def test_the_daily_report_says_what_the_meat_earned_and_what_to_set_aside(self):
        """El parte tenía los kilos; ahora dice si han valido la pena.

        Kilos y unidades dicen cuánto se ha movido, no cuánto se ha ganado. El
        parte del día trae ahora lo que se ingresó con la carne que salió, lo
        que costó esa carne, y —si la casa ha dicho cuánto paga— lo que hay que
        apartar y lo que queda limpio.

        Se monta con el banco de pruebas, que es un mes de trabajo de verdad
        con sus ventas: un ejemplo hecho a mano probaría la suma, no el parte.
        """
        import pathlib
        import tempfile
        from datetime import date as _date

        from thegrill import bench
        from thegrill.meat import service as meat_service
        from thegrill.models import Restaurant
        from thegrill.web import impuestos

        carpeta = pathlib.Path(tempfile.mkdtemp())
        db.init_engine(f"sqlite:///{carpeta/'parte.db'}")
        db.create_all()
        hasta = _date(2026, 9, 20)
        with db.session_scope() as s:
            casa = bench.build(s, days=20, seed=3, multisite=False, index=90, until=hasta)
            rid = casa.restaurant_id

        with db.session_scope() as s:
            # El día con más movimiento de los últimos, que es donde hay ventas.
            partes = [meat_service.daily_report(s, rid, on=hasta - timedelta(days=d))
                      for d in range(6)]
            parte = max(partes, key=lambda p: p.sales_revenue)
            assert parte.sales_revenue > 0, "ningún día del banco vendió nada"
            assert parte.sales_cost > 0
            assert parte.sales_margin == round(parte.sales_revenue - parte.sales_cost, 2)
            # Y el food cost del día sale de esos mismos dos números.
            assert parte.sales_food_cost_pct == round(
                parte.sales_cost / parte.sales_revenue * 100, 2)

            # Sin tipo puesto no se aparta nada.
            restaurante = s.get(Restaurant, rid)
            assert not impuestos.de_la_casa(restaurante, parte.sales_margin).hay_impuesto

            # Con el 27 puesto, el margen se parte y las dos partes suman.
            restaurante.tax_pct = 27.0
            corte = impuestos.de_la_casa(restaurante, parte.sales_margin)
            assert corte.hay_impuesto
            assert corte.impuesto == round(parte.sales_margin * 0.27, 2)
            assert round(corte.impuesto + corte.limpio, 2) == round(parte.sales_margin, 2)

    def test_the_house_says_what_it_pays_in_tax_and_the_margin_gets_split(self, client):
        """El margen bruto no es lo que queda: una parte se la lleva Hacienda.

        Un lomo que se compra a 300 y se vende a 900 deja 600 de margen, pero de
        esos 600 no se lleva la casa 600. Un restaurante que mira el bruto y
        gasta contra él va bien once meses y mal el doceavo, siempre el mismo.
        El manager dice cuánto paga y el programa parte el margen en dos: lo que
        hay que apartar y lo que queda.
        """
        from tests.meat_helpers import csrf_from

        signup(client)
        pantalla = client.get("/configuracion")
        assert 'name="tax_pct"' in pantalla.text
        assert client.post("/configuracion", data={
            "csrf": csrf_from(pantalla.text), "language": "es",
            "tax_pct": "27"}).status_code == 303
        with db.session_scope() as s:
            from thegrill.models import Restaurant
            casa = s.query(Restaurant).filter(Restaurant.platform.isnot(True)).one()
            assert casa.tax_pct == 27.0
        # Y vuelve escrito en la pantalla, que es lo que dice que se guardó.
        assert 'value="27"' in client.get("/configuracion").text

    def test_without_a_rate_nothing_is_split_and_nothing_is_invented(self, client):
        """Sin decir cuánto paga no se enseña ningún reparto.

        Poner el tipo del país, o una media, sería pintar un número que parece
        de la casa y no lo es, y sobre él se toman decisiones.
        """
        from thegrill.web import impuestos

        signup(client)
        with db.session_scope() as s:
            from thegrill.models import Restaurant
            casa = s.query(Restaurant).filter(Restaurant.platform.isnot(True)).one()
            assert casa.tax_pct is None
            assert impuestos.tipo_de(casa) == 0.0
            assert not impuestos.de_la_casa(casa, 600.0).hay_impuesto
            assert impuestos.de_la_casa(casa, 600.0).limpio == 600.0

    def test_what_it_cost_to_bring_it_is_shared_by_every_kilo_of_the_lorry(self, client):
        """La carne no cuesta lo que dice la factura: cuesta ponerla en la cámara.

        El precio del kilo es de cada pieza —dos bolsas del mismo camión no
        valen lo mismo si una es MB9 y la otra MB6— y por eso ya no hay un
        precio del lote que se copie a todas. Lo que sí es del camión entero es
        lo que costó traerlo: el flete y la aduana, que solo hay cuando viene
        de fuera y se reparten a cada kilo que traía.

        Un lomo a 40 con 2 de transporte y 1 de aduana no está a 40: está a 43,
        y despiezarlo como si estuviera a 40 se lleva ese siete por ciento a
        todos los cortes, a la carta y al food cost del mes sin que nadie lo
        vea. Y los tres números se guardan por separado, que es lo que contesta
        «¿me sale más caro el australiano por la carne o por el flete?».
        """
        signup(client)
        self.recibir(client, **{"price:0": "40", "price:1": "40",
                                "freight_kg": "2", "duty_kg": "1"})
        with db.session_scope() as s:
            piezas = {p.serial: p for p in s.query(Primal)}
            una = piezas["8017"]
            assert una.landed_usd_per_kg == 43.0          # puesto en la cámara
            assert una.goods_usd_per_kg == 40.0           # y de qué está hecho
            assert una.freight_usd_per_kg == 2.0
            assert una.duty_usd_per_kg == 1.0
            assert una.piece_cost_usd == round(9.4 * 43.0, 2)
            # A las dos bolsas del mismo camión les toca el mismo flete.
            assert piezas["8018"].landed_usd_per_kg == 43.0

    def test_a_lorry_from_here_does_not_pay_freight_or_duty(self, client):
        """Sin importación no hay nada que sumar, y el kilo es el de la factura."""
        signup(client)
        self.recibir(client, **{"price:0": "40", "price:1": "40"})
        with db.session_scope() as s:
            una = s.query(Primal).filter_by(serial="8017").one()
            assert una.landed_usd_per_kg == 40.0
            assert una.freight_usd_per_kg is None and una.duty_usd_per_kg is None

    def test_freight_alone_never_invents_a_price_for_a_piece_that_has_none(self, client):
        """Sumarle el flete a lo que no se sabe lo que cuesta sería inventarlo.

        En el muelle casi nunca se sabe el precio: llega con la factura, días
        después. Una pieza así tiene que quedarse esperando, no entrar valiendo
        dos euros el kilo porque esos dos son lo que costó el camión.
        """
        signup(client)
        self.recibir(client, **{"freight_kg": "2", "duty_kg": "1"})
        with db.session_scope() as s:
            una = s.query(Primal).filter_by(serial="8017").one()
            assert una.landed_usd_per_kg is None, "ha entrado valiendo solo el flete"
            assert una.freight_usd_per_kg == 2.0        # pero apuntado queda
        assert "8017" in client.get("/recepcion/precios").text

    def test_the_price_set_later_still_pays_the_freight(self, client):
        """Y cuando dirección pone el precio, lo de traerla sigue siendo suyo."""
        from tests.meat_helpers import csrf_from

        signup(client)
        self.recibir(client, **{"freight_kg": "2", "duty_kg": "1"})
        pantalla = client.get("/recepcion/precios")
        assert client.post("/recepcion/precios",
                           data={"csrf": csrf_from(pantalla.text), "serial": ["8017"],
                                 "price:8017": "40"}).status_code in (200, 303)
        with db.session_scope() as s:
            una = s.query(Primal).filter_by(serial="8017").one()
            assert una.goods_usd_per_kg == 40.0
            assert round(una.landed_usd_per_kg, 6) == 43.0

    def test_the_manager_who_receives_prices_on_the_spot(self, client):
        """Quien ve el dinero no necesita el segundo paso."""
        signup(client)
        pantalla = client.get("/recepcion").text
        assert 'name="price:0"' in pantalla
        # El precio del kilo es de cada pieza y no del camión: del camión solo
        # es lo que costó traerla, que se reparte a todos sus kilos.
        assert 'name="price_kg"' not in pantalla
        assert 'name="freight_kg"' in pantalla and 'name="duty_kg"' in pantalla
        self.recibir(client, **{"price:0": "32", "price:1": "32"})
        with db.session_scope() as s:
            assert all(p.landed_usd_per_kg == 32 for p in s.query(Primal))
            assert s.query(Primal).count() == 2
        assert "8017" in client.get("/despiece").text
        # Y no hay nada esperando ni nadie a quien avisar.
        assert t_sin_precio(client) == 0

    def test_what_is_waiting_shows_up_on_the_front_screen(self, client):
        signup(client)
        paco = add_user(client, email="paco@marina.com", name="Paco", role=Role.BUTCHER)
        self.recibir(paco)
        assert t_sin_precio(client) == 2
        assert "2" in client.get("/hoy").text

    # ------------------------------------------------------------- la foto
    def test_a_piece_carries_the_photo_of_its_label(self, client, tmp_path, monkeypatch):
        monkeypatch.setattr(meatapp, "UPLOAD_DIR", str(tmp_path / "subidas"))
        signup(client)
        self.recibir(client, **{"price:0": "32"})
        token = csrf_from(client.get("/recepcion").text)

        r = client.post("/carne/8017/foto", data={"csrf": token, "next": "/recepcion"},
                        files={"foto": ("etiqueta.png", PNG, "image/png")})
        assert r.status_code == 303
        with db.session_scope() as s:
            guardada = s.query(Primal).filter_by(serial="8017").one().photo_ref
            assert guardada and os.path.exists(guardada)

        # Se sirve, pero solo a gente de esta casa y no desde /static.
        foto = client.get("/carne/8017/etiqueta")
        assert foto.status_code == 200 and foto.content == PNG
        assert "/static/" not in client.get("/trazabilidad?serial=8017").text.split(
            "etiqueta")[0][-120:]

    def test_a_second_photo_replaces_the_first_and_the_old_file_goes(self, client,
                                                                     tmp_path, monkeypatch):
        monkeypatch.setattr(meatapp, "UPLOAD_DIR", str(tmp_path / "subidas"))
        signup(client)
        self.recibir(client, **{"price:0": "32"})
        token = csrf_from(client.get("/recepcion").text)
        client.post("/carne/8017/foto", data={"csrf": token},
                    files={"foto": ("uno.png", PNG, "image/png")})
        with db.session_scope() as s:
            primera = s.query(Primal).filter_by(serial="8017").one().photo_ref
        client.post("/carne/8017/foto", data={"csrf": token},
                    files={"foto": ("dos.png", PNG, "image/png")})
        with db.session_scope() as s:
            segunda = s.query(Primal).filter_by(serial="8017").one().photo_ref
        assert segunda != primera
        assert os.path.exists(segunda) and not os.path.exists(primera)

    def test_a_file_that_is_not_a_photo_is_refused(self, client, tmp_path, monkeypatch):
        monkeypatch.setattr(meatapp, "UPLOAD_DIR", str(tmp_path / "subidas"))
        signup(client)
        self.recibir(client, **{"price:0": "32"})
        token = csrf_from(client.get("/recepcion").text)
        r = client.post("/carne/8017/foto", data={"csrf": token},
                        files={"foto": ("virus.exe", b"MZ", "application/x-msdownload")})
        assert r.status_code == 303
        with db.session_scope() as s:
            assert s.query(Primal).filter_by(serial="8017").one().photo_ref is None


def t_sin_precio(client) -> int:
    from thegrill.meat import service as meatsvc
    with db.session_scope() as s:
        casa = s.query(Restaurant).filter(Restaurant.platform.isnot(True)).one()
        return len(meatsvc.awaiting_price(s, casa.id))


# Un PNG de un píxel: lo justo para que el navegador y el programa lo den por foto.
PNG = (b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01\x08\x06"
       b"\x00\x00\x00\x1f\x15\xc4\x89\x00\x00\x00\nIDATx\x9cc\x00\x01\x00\x00\x05"
       b"\x00\x01\r\n-\xb4\x00\x00\x00\x00IEND\xaeB`\x82")


class TestRecepcionDeUnaEnUna:
    """Se descarga pieza a pieza, y la foto se hace al coger la bolsa.

    La pantalla de antes abría con ocho líneas y un botón de añadir ocho más.
    En el muelle no se descarga así: se coge una bolsa, se le mira la etiqueta,
    se apunta y se coge la siguiente. Y la foto es de la etiqueta que tienes
    delante, no de una de ocho.
    """

    def test_the_screen_asks_for_one_piece_and_the_camera_first(self, client):
        signup(client)
        pantalla = client.get("/recepcion").text
        assert 'name="serial:0"' in pantalla
        assert 'name="serial:1"' not in pantalla        # una cada vez
        assert "masfilas" not in pantalla               # sin «añadir 8 líneas»
        # La foto va en el propio formulario y abre la cámara de atrás.
        assert 'enctype="multipart/form-data"' in pantalla
        assert 'capture="environment"' in pantalla
        # El orden: primero se configura el lote, luego la foto de la etiqueta
        # de la bolsa que tienes en la mano, y al final sus números.
        assert (pantalla.index('name="lot"') < pantalla.index('name="foto"')
                < pantalla.index('name="g:0"'))

    def test_the_lot_says_which_one_it_is_once_it_is_set(self, client):
        """Se configura una vez y se le van subiendo piezas: hay que ver cuál es."""
        signup(client)
        form = client.get("/recepcion")
        r = client.post("/recepcion", data={
            "csrf": csrf_from(form.text), "lot": "L-QUIEN", "sku": "Ribeye AUS MB7",
            "price:0": "32", "serial:0": "9500", "g:0": "9100"})
        assert r.status_code == 303
        pantalla = client.get("/recepcion").text
        assert "L-QUIEN" in pantalla and "Ribeye AUS MB7" in pantalla
        # Y plegado, porque ya está configurado: lo que se toca es la pieza.
        assert "<details" in pantalla and 'name="lot"' in pantalla

    def test_the_label_travels_with_the_piece_in_one_go(self, client, tmp_path,
                                                        monkeypatch):
        monkeypatch.setattr(meatapp, "UPLOAD_DIR", str(tmp_path / "subidas"))
        signup(client)
        form = client.get("/recepcion")
        r = client.post("/recepcion", data={
            "csrf": csrf_from(form.text), "lot": "L-UNA", "sku": "Ribeye AUS",
            "producer_plant": "Teys Biloela", "serial:0": "9200", "g:0": "9200",
            "price:0": "32"},
            files={"foto": ("etiqueta.png", PNG, "image/png")})
        assert r.status_code == 303
        with db.session_scope() as s:
            pieza = s.query(Primal).filter_by(serial="9200").one()
            assert pieza.producer_plant == "Teys Biloela"
            assert pieza.photo_ref and os.path.exists(pieza.photo_ref)

    def test_what_belongs_to_the_delivery_stays_typed_for_the_next_piece(self, client):
        """Veinte piezas de la misma caja no son veinte veces el matadero."""
        signup(client)
        form = client.get("/recepcion")
        r = client.post("/recepcion", data={
            "csrf": csrf_from(form.text), "lot": "L-CAMION", "sku": "Ribeye AUS",
            "grade": "MB7", "origin": "AUS", "producer_plant": "Teys Biloela",
            "est_code": "AUS 1234", "breed": "Angus", "price:0": "32",
            "serial:0": "9201", "g:0": "9200"})
        # Guardar redirige, y lo del camión llega con la pantalla de detrás:
        # es la misma caja, así que no se vuelve a teclear el matadero.
        assert r.status_code == 303
        pantalla = client.get("/recepcion").text
        for valor in ("L-CAMION", "Ribeye AUS", "Teys Biloela", "AUS 1234", "Angus"):
            assert f'value="{valor}"' in pantalla, valor
        # Y lo de la pieza se vacía: la siguiente es otra bolsa. La casilla de
        # los kilos vuelve en blanco y con el cursor puesto, que es donde va a
        # escribir el de fuera. Se mira lo que hace la casilla, no cómo está
        # escrita: el `autofocus` va detrás de una condición y no pegado.
        import re
        casilla = re.search(r'<input name="g:0"[^>]*>', pantalla, re.S)
        assert casilla, "no está la casilla del peso"
        assert "autofocus" in casilla.group(0)
        assert 'value=""' in casilla.group(0)

    def test_the_screen_says_to_write_the_number_on_the_meat(self, client):
        """El número tiene que estar encima de la carne, no solo en la base.

        En la cámara nadie abre el móvil para saber qué bolsa tiene en la mano:
        si la pieza no lleva su número escrito, la trazabilidad vive en el
        ordenador y en la cámara se busca a ojo.
        """
        signup(client)
        pantalla = client.get("/recepcion").text
        assert "rotulador permanente" in pantalla and "etiqueta pegada" in pantalla

        # Y al guardar se dice con el número delante, que es cuando se coge el
        # rotulador: la pieza todavía está en la mano.
        form = client.get("/recepcion")
        r = client.post("/recepcion", data={
            "csrf": csrf_from(form.text), "lot": "L-ROTU", "sku": "Ribeye AUS",
            "price:0": "30", "serial:0": "9210", "g:0": "9100"})
        # Con el número y el peso, que es lo que hay que escribir encima. El
        # recado llega con la pantalla de detrás, no como respuesta al POST.
        assert r.status_code == 303
        recado = client.get("/recepcion").text
        assert "9210 · 9,1 kg" in recado and "en el primal" in recado

    def test_a_phone_with_no_signal_still_books_the_piece(self, client):
        """La foto necesita línea; lo escrito, no. Lo escrito manda."""
        signup(client)
        form = client.get("/recepcion")
        # Sin fichero: es lo que manda la cola del teléfono cuando vuelve.
        r = client.post("/recepcion", data={
            "csrf": csrf_from(form.text), "lot": "L-SINRED", "sku": "Ribeye AUS",
            "price:0": "30", "serial:0": "9202", "g:0": "8800",
            "envio": "numero-de-la-cola"})
        assert r.status_code == 303
        with db.session_scope() as s:
            pieza = s.query(Primal).filter_by(serial="9202").one()
            assert pieza.photo_ref is None      # se hace luego, desde la lista
            assert pieza.weight_kg == 8.8


def test_the_butchery_takes_the_weight_of_the_whole_tray(client):
    """En la mesa se pesa la bandeja, no filete a filete.

    Se escriben las piezas que han salido y los kilos de todas juntas, y el
    peso de cada una sale de ahí. Pedir los gramos de una pieza obliga a pesar
    una y fiarse, o a hacer la división a mano con las manos llenas de grasa.
    """
    signup(client)
    items = setup_cuts(client)
    deliver(client)
    form = client.get("/despiece")
    r = client.post("/despiece", data={
        "csrf": csrf_from(form.text), "tg": "TG-0001", "date": str(HOY),
        "before_g": "10000", "waste_g": "600", "primal": "8017",
        "cut:0": "Striploin steak", "item:0": items["Striploin steak"],
        "pieces:0": "18", "total:0": "5400", "index:0": "1"})
    assert r.status_code == 303
    with db.session_scope() as s:
        corte = s.query(DespieceCut).filter_by(cut_name="Striploin steak").one()
        assert corte.pieces == 18
        assert corte.weight_per_piece_g == 300      # 5,4 kg entre 18
        assert corte.total_kg == pytest.approx(5.4)


def test_grams_per_piece_still_work_for_what_comes_from_the_paper(client):
    """La hoja de papel los pide así, y la cola de un teléfono puede traerlos."""
    signup(client)
    items = setup_cuts(client)
    deliver(client)
    form = client.get("/despiece")
    r = client.post("/despiece", data={
        "csrf": csrf_from(form.text), "tg": "TG-0002", "date": str(HOY),
        "before_g": "10000", "waste_g": "600", "primal": "8017",
        "cut:0": "Striploin steak", "item:0": items["Striploin steak"],
        "pieces:0": "20", "grams:0": "250", "index:0": "1"})
    assert r.status_code == 303
    with db.session_scope() as s:
        corte = s.query(DespieceCut).filter_by(cut_name="Striploin steak").one()
        assert corte.weight_per_piece_g == 250


class TestComoLlega:
    """Cómo bajó del camión, que no es lo mismo que dónde está ahora.

    Una pieza que llega congelada y una que llega fresca y se mete al arcón
    acaban las dos en el congelador, pero no son la misma carne: la primera
    nunca estuvo fresca en esta casa. Y la temperatura de la caja al abrirla es
    lo primero que se pregunta el día que una pieza sale mal.
    """

    def recibir(self, client, **extra):
        from thegrill.models import Storage
        form = client.get("/recepcion")
        data = {"csrf": csrf_from(form.text), "lot": "L-FRIO", "sku": "Ribeye AUS",
                "price:0": "32", "use_by": str(HOY + timedelta(days=40)),
                "serial:0": "9400", "g:0": "9400"}
        data.update(extra)
        return client.post("/recepcion", data=data)

    def test_chilled_meat_stays_in_the_chiller(self, client):
        from thegrill.models import Storage
        signup(client)
        self.recibir(client, arrival="CHILLED", arrival_c="2,4")
        with db.session_scope() as s:
            pieza = s.query(Primal).filter_by(serial="9400").one()
            assert pieza.arrival == Storage.CHILLED
            assert pieza.arrival_c == 2.4
            assert pieza.storage == Storage.CHILLED
            assert not pieza.frozen_on_arrival

    def test_meat_that_arrives_frozen_is_in_the_freezer_from_day_one(self, client):
        """Contarla como fresca le pondría el reloj que no es."""
        from thegrill.models import Storage
        signup(client)
        self.recibir(client, arrival="FROZEN", arrival_c="-19")
        with db.session_scope() as s:
            pieza = s.query(Primal).filter_by(serial="9400").one()
            assert pieza.arrival == Storage.FROZEN
            assert pieza.arrival_c == -19
            assert pieza.storage == Storage.FROZEN
            assert pieza.storage_since == HOY

    def test_fresh_meat_that_goes_straight_to_the_freezer(self, client):
        """Llega fresca y no pasa por la cámara: el arcón manda desde hoy."""
        from thegrill.models import Storage
        signup(client)
        self.recibir(client, arrival="CHILLED", arrival_c="1,8", frozen_on_arrival="1")
        with db.session_scope() as s:
            pieza = s.query(Primal).filter_by(serial="9400").one()
            assert pieza.arrival == Storage.CHILLED     # llegó fresca, eso no cambia
            assert pieza.frozen_on_arrival is True
            assert pieza.storage == Storage.FROZEN      # pero está en el arcón

    def test_what_already_comes_frozen_is_not_frozen_again(self, client):
        """La casilla solo tiene sentido si llega fresca."""
        signup(client)
        self.recibir(client, arrival="FROZEN", frozen_on_arrival="1")
        with db.session_scope() as s:
            assert not s.query(Primal).filter_by(serial="9400").one().frozen_on_arrival

    def test_the_piece_sheet_says_how_it_arrived(self, client):
        signup(client)
        self.recibir(client, arrival="CHILLED", arrival_c="2,4", frozen_on_arrival="1")
        ficha = client.get("/trazabilidad?serial=9400").text
        assert "Cómo llegó" in ficha
        assert "Refrigerada a 2,4 °C" in ficha
        assert "Congelada al entrar" in ficha
