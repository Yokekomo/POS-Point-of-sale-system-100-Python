"""El parte de ventas del POS, leído de un fichero.

Cada punto de venta exporta lo suyo, pero todos exportan lo mismo: una línea
por artículo con lo que se ha vendido. Cambian los nombres de las columnas, el
separador y la manera de escribir los números —«1.234,56» aquí, «1,234.56»
allí—, y por eso esto no pide un formato concreto: lee el fichero, mira la
cabecera y dice qué ha entendido de cada columna.

Lo que se busca en la cabecera:

- el **código** del artículo y su **nombre**: con uno basta, y si vienen los
  dos manda el código, porque un nombre se reescribe y un código no;
- las **unidades** vendidas;
- el **peso**, en kilos o en gramos, para lo que se cobra a peso;
- el **importe**, que no hace falta para descontar pero se enseña para cuadrar.

Leer no descuenta nada. Primero se enseña lo entendido —qué producto es cada
línea, cuáles no se reconocen— y solo después, si se confirma, se descuenta.
Un parte mal leído que descuenta a ciegas es peor que no tener parte.
"""
import csv
import io
import re
import unicodedata
from dataclasses import dataclass, field

MAX_BYTES = 4 * 1024 * 1024     # un parte de ventas no pesa más que esto
MAX_ROWS = 5000                 # ni trae más líneas que esto
EXTENSIONS = (".csv", ".tsv", ".txt", ".xlsx", ".xlsm")

# Cómo llama cada POS a cada columna. Se compara sin tildes ni mayúsculas.
HEADERS = {
    "code": ("codigo", "cod", "code", "plu", "sku", "referencia", "ref", "id",
             "articulo", "art", "item", "item code", "product code", "codigo articulo"),
    "name": ("nombre", "producto", "descripcion", "description", "product", "name",
             "articulo", "item", "item name", "product name", "concepto", "plato"),
    "units": ("unidades", "uds", "ud", "cantidad", "cant", "qty", "quantity", "units",
              "count", "n", "num", "vendidos", "sold"),
    "kg": ("kg", "kgs", "kilos", "kilogramos", "peso", "peso kg", "weight", "weight kg"),
    "grams": ("g", "gr", "gramos", "grams", "peso g", "weight g"),
    "amount": ("importe", "total", "amount", "venta", "ventas", "neto", "net", "pvp",
               "revenue", "sales", "importe total", "total venta"),
}


class ImportError_(ValueError):
    """El fichero no se puede leer, y se dice por qué."""


@dataclass
class SaleRow:
    """Una línea del parte, ya entendida."""
    line: int                       # la fila del fichero, para poder señalarla
    code: str | None = None
    name: str | None = None
    units: float = 0.0
    kg: float | None = None
    amount: float | None = None

    @property
    def key(self) -> str:
        """Con qué se busca en el mapeo del POS: el código si lo hay."""
        return (self.code or self.name or "").strip()


@dataclass
class Parsed:
    rows: list[SaleRow] = field(default_factory=list)
    columns: dict[str, str] = field(default_factory=dict)   # qué es cada columna
    warnings: list[str] = field(default_factory=list)
    skipped: int = 0                # filas sin unidades o sin artículo

    @property
    def units(self) -> float:
        """Las unidades vendidas que trae el parte, todas sumadas."""
        return round(sum(r.units for r in self.rows), 4)

    @property
    def kg(self) -> float:
        """Los kilos del parte, para lo que se vende al peso."""
        return round(sum(r.kg or 0.0 for r in self.rows), 4)

    @property
    def amount(self) -> float:
        """Lo facturado según el parte."""
        return round(sum(r.amount or 0.0 for r in self.rows), 2)


def parse(data: bytes, filename: str = "") -> Parsed:
    """Lee el fichero y devuelve lo que ha entendido. No toca el almacén."""
    if not data:
        raise ImportError_("El fichero está vacío")
    if len(data) > MAX_BYTES:
        raise ImportError_(f"El fichero pasa de {MAX_BYTES // (1024 * 1024)} MB")
    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm")):
        table = _from_excel(data)
    else:
        table = _from_text(data)
    if not table:
        raise ImportError_("El fichero no tiene ninguna línea")
    return _understand(table)


# Una celda es texto —y entonces hay que averiguar cómo escribe los decimales—
# o es un número que ya venía leído de Excel, y entonces no hay nada que
# averiguar.
Celda = str | float


# ------------------------------------------------------------------ lectura
def _from_text(data: bytes) -> list[list[str]]:
    """Lee un fichero de texto separado por comas, puntos y comas o tabuladores.

    Cada caja escribe el suyo a su manera, así que el separador se adivina
    mirando las primeras líneas; si no hay quien lo adivine, gana el que más
    veces sale en la cabecera. Las filas vacías se tiran.
    """
    text = _decode(data)
    sample = "\n".join(text.splitlines()[:20])
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t|")
        delimiter = dialect.delimiter
    except csv.Error:
        # Sin pistas, gana el separador que más veces aparezca en la cabecera.
        head = text.splitlines()[0] if text.splitlines() else ""
        delimiter = max(";,\t|", key=head.count) if head else ","
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    return [[cell.strip() for cell in row] for row in reader if any(c.strip() for c in row)]


def _decode(data: bytes) -> str:
    """Averigua con qué juego de letras está escrito el fichero.

    Se prueban de más moderno a más viejo. Sin esto, un parte guardado en
    Windows llega con la eñe rota y el artículo no casa con el de la carta.
    """
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ImportError_("No se entiende cómo está escrito el fichero")


def _from_excel(data: bytes) -> list[list[Celda]]:
    """Lo que Excel ya leyó como número no se vuelve a interpretar.

    Pasarlo por `str()` lo convierte en «12.5», y si el resto del parte
    escribe los decimales con coma, ese punto se lee después como separador de
    miles: 12,5 kg entran como 125 y el escandallo del mes queda multiplicado
    por diez sin que nadie vea nada raro. Un número que viene de una celda
    numérica ya no tiene ninguna duda que resolver, así que viaja como número
    hasta el final.
    """
    try:
        from openpyxl import load_workbook
    except ImportError:                                   # pragma: no cover
        raise ImportError_("Este servidor no sabe leer Excel; exporta el parte en CSV") from None
    try:
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        raise ImportError_("Ese Excel no se puede abrir; exporta el parte en CSV") from None
    sheet = book.worksheets[0]
    table = []
    for row in sheet.iter_rows(values_only=True):
        cells: list[Celda] = []
        for c in row:
            if c is None or isinstance(c, bool):
                cells.append("" if c is None else str(c))
            elif isinstance(c, (int, float)):
                cells.append(float(c))        # ya es un número: no se toca
            else:
                cells.append(str(c).strip())
        if any(c != "" for c in cells):
            table.append(cells)
        if len(table) > MAX_ROWS + 5:
            break
    book.close()
    return table


# --------------------------------------------------------------- lo que dice
def _understand(table: list[list[Celda]]) -> Parsed:
    """Convierte la tabla en ventas: qué columna es cada cosa y qué dice cada fila.

    Dos cosas hay que acertar antes de leer nada: dónde está la cabecera —los
    partes traen encima el nombre del local y la fecha— y cómo escribe los
    decimales este fichero. Lo segundo se decide mirando el parte entero: en
    uno donde «1,236» es un kilo y pico, no puede ser mil en la fila de abajo.
    """
    out = Parsed()
    header, start = _find_header(table)
    if header is None:
        raise ImportError_(
            "No se encuentra la cabecera. El parte tiene que traer una fila con los "
            "nombres de las columnas: el artículo y las unidades, como mínimo.")

    mapping = _map_columns(header)
    if "units" not in mapping:
        raise ImportError_("No hay columna de unidades vendidas")
    if "code" not in mapping and "name" not in mapping:
        raise ImportError_("No hay columna de artículo: ni código ni nombre")
    out.columns = {field: str(header[index]) for field, index in mapping.items()}

    # Cómo escribe los decimales este fichero. Se decide mirándolo entero y no
    # celda a celda: «1,236» es un kilo y pico o son mil, pero dentro del mismo
    # parte no son las dos cosas.
    decimal, clear = _decimal_separator(_numeric_cells(table[start:], mapping))
    if not clear:
        ejemplo = next((c for c in _numeric_cells(table[start:], mapping)
                        if "," in c or "." in c), None)
        if ejemplo:
            out.warnings.append(
                f"No está claro cómo escribe los decimales: «{ejemplo}» se ha leído como "
                f"{number_of(ejemplo, decimal):.10g}. Si no es eso, exporta el parte con "
                "punto decimal.")

    for number, row in enumerate(table[start:], start=start + 1):
        if len(out.rows) >= MAX_ROWS:
            out.warnings.append(f"Solo se han leído las primeras {MAX_ROWS} líneas")
            break
        parsed = _row(row, mapping, number, decimal)
        if parsed is None:
            out.skipped += 1
            continue
        out.rows.append(parsed)

    if not out.rows:
        raise ImportError_("El fichero no trae ninguna venta con unidades")
    return out


def _find_header(table: list[list[Celda]]) -> tuple[list[Celda] | None, int]:
    """La cabecera es la primera fila que nombra al menos dos cosas conocidas."""
    for index, row in enumerate(table[:10]):
        found = _map_columns(row)
        if len(found) >= 2 and ("code" in found or "name" in found):
            return row, index + 1
    return None, 0


def _map_columns(header: list[Celda]) -> dict[str, int]:
    """Qué columna es cada cosa, por el nombre de la cabecera.

    Cada caja llama a las cosas a su manera —«artículo», «producto», «plu»— y
    aquí se traducen todas al mismo sitio. La primera que casa se queda con el
    puesto: un parte con dos columnas parecidas no se queda sin ninguna.
    """
    mapping: dict[str, int] = {}
    for index, cell in enumerate(header):
        key = _norm(cell)
        if not key:
            continue
        for field_name, names in HEADERS.items():
            if field_name in mapping:
                continue
            if key in names:
                mapping[field_name] = index
                break
    return mapping


def _row(row: list[Celda], mapping: dict[str, int], number: int,
         decimal: str = ".") -> SaleRow | None:
    """Una fila del parte convertida en venta, o nada si no lo es.

    Se cae lo que no es una venta: los totales, los subtotales y las filas sin
    artículo. Vender cero o menos tampoco es vender.
    """
    def cell(field_name: str) -> Celda:
        """El valor de esa columna en esta fila, o vacío si no está."""
        index = mapping.get(field_name)
        if index is None or index >= len(row):
            return ""
        valor = row[index]
        return valor.strip() if isinstance(valor, str) else valor

    def texto(field_name: str) -> str:
        """Ese valor como texto, sin que un código numérico salga con decimales.

        Una hoja de cálculo devuelve el código 4070 como 4070.0, y así no casa con
        ningún artículo de la carta.
        """
        valor = cell(field_name)
        if isinstance(valor, float):        # un código de artículo numérico
            return f"{valor:.10g}"
        return str(valor)

    units = number_of(cell("units"), decimal)
    if units is None or units <= 0:
        return None
    code, name = texto("code"), texto("name")
    if not code and not name:
        return None

    kg = number_of(cell("kg"), decimal)
    if kg is None:
        grams = number_of(cell("grams"), decimal)
        kg = round(grams / 1000, 6) if grams else None
    return SaleRow(line=number, code=code or None, name=name or None, units=units,
                   kg=kg if kg and kg > 0 else None,
                   amount=number_of(cell("amount"), decimal))


def _numeric_cells(rows: list[list[Celda]], mapping: dict[str, int]) -> list[str]:
    """Las celdas de texto que llevan números: las únicas que dicen cómo se
    escriben los decimales en este fichero. Las que ya venían como número no
    opinan, porque no tienen separador que interpretar."""
    columnas = [mapping[f] for f in ("units", "kg", "grams", "amount") if f in mapping]
    out = []
    for row in rows[:200]:
        for index in columnas:
            if index < len(row) and isinstance(row[index], str) and row[index].strip():
                out.append(row[index].strip())
    return out


def _decimal_separator(cells: list[str]) -> tuple[str, bool]:
    """Qué separador usa este fichero para los decimales, y si está claro.

    Con los dos separadores en la misma celda no hay duda: el último es el
    decimal. Con uno solo lo delatan dos cosas: el cero delante —nadie separa
    los miles detrás de un cero, así que «0,824» es menos de un kilo— y el
    grupo que no tiene tres cifras, porque quien escribe «12,5» usa la coma de
    decimal. Si todos los grupos son de tres cifras y ninguno empieza por cero
    no hay manera de saberlo: se toma el punto y se avisa.
    """
    for text in cells:
        if "," in text and "." in text:
            return ("," if text.rfind(",") > text.rfind(".") else "."), True
    for text in cells:
        # Nadie separa los miles detrás de un cero: «0,824» es menos de un kilo.
        cero = re.match(r"^\s*-?0([.,])\d+\s*$", text)
        if cero:
            return cero.group(1), True
    for text in cells:
        for trozo in re.findall(r",(\d+)", text):
            if len(trozo) != 3:
                return ",", True
    for text in cells:
        for trozo in re.findall(r"\.(\d+)", text):
            if len(trozo) != 3:
                return ".", True
    # Todos los grupos de tres cifras y ninguno detrás de un cero: «1.234» son
    # mil doscientos treinta y cuatro o son uno y pico, y el fichero no lo
    # dice. Antes se tomaba el punto por decimal y se daba por seguro, que es
    # la peor de las dos opciones: el parte entraba dividido entre mil y no
    # había manera de enterarse. Se sigue tomando el punto, pero se avisa.
    return ".", not any(("," in c or "." in c) for c in cells)


# ------------------------------------------------------------------ números
def number_of(raw: str | None, decimal: str = ".") -> float | None:
    """Un número escrito como lo escriba el POS: 1.234,56 o 1,234.56 o 12 kg.

    `decimal` es el separador que usa ese fichero, decidido mirándolo entero.
    """
    if raw is None:
        return None
    if isinstance(raw, (int, float)) and not isinstance(raw, bool):
        return float(raw)                               # venía de una celda numérica
    text = re.sub(r"[^\d,.\-]", "", str(raw).strip())   # fuera monedas y unidades
    if not text or text in ("-", ".", ","):
        return None
    miles = "." if decimal == "," else ","
    text = text.replace(miles, "").replace(decimal, ".")
    try:
        return float(text)
    except ValueError:
        return None


def _norm(value: str) -> str:
    """Sin tildes, sin mayúsculas y sin dobles espacios, para comparar cabeceras."""
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(c for c in text if not unicodedata.combining(c))
    text = re.sub(r"[^\w\s]", " ", text.lower())
    return " ".join(text.split())
