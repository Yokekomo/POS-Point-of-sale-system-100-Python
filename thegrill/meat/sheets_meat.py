"""Hojas de carne en Excel, listas para imprimir y rellenar a mano.

Son el respaldo de papel de la cámara: se imprimen, se cuelgan al lado de la
balanza y se rellenan con bolígrafo cuando no hay móvil a mano. Las columnas
del papel son las mismas que pide la pantalla, así que pasarlo luego es copiar
en orden.

No llevan fórmulas: son rejillas para escribir a mano.
"""
import io
from dataclasses import dataclass, field
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from thegrill.models import Restaurant
from thegrill.web.i18n import DEFAULT_LANG, direction, t
from thegrill.web.sheets import (BLANK_ROWS, EXAMPLE_FILL, FONT, GRID, HEADER_FILL, INK,
                                 MUTED, PORTRAIT_MAX_WIDTH, WRITE_LINE, filename_for,
                                 sheet_title)


@dataclass(frozen=True)
class Column:
    key: str             # clave de texto para el encabezado
    example: str         # cómo se rellena, en gris
    width: int = 16


@dataclass(frozen=True)
class Sheet:
    code: str
    title_key: str
    sub_key: str
    columns: list[Column] = field(default_factory=list)


SHEETS: list[Sheet] = [
    Sheet("recepcion", "m.rec.title", "m.rec.sub", [
        Column("m.rec.lot", "DXB20260910", 18),
        Column("m.rec.serial", "8017", 12),
        Column("m.rec.sku", "Striploin AUS", 24),
        Column("m.rec.grade", "MB9+", 10),
        Column("m.rec.origin", "AUS", 10),
        Column("m.rec.kg", "9,4", 10),
        Column("m.rec.price_kg", "32,00", 12),
        Column("m.rec.use_by", "2026-12-20", 14),
    ]),
    Sheet("despiece", "m.tg.title", "m.tg.sub", [
        Column("m.tg.number", "TG-0010", 12),
        Column("m.rec.serial", "8017", 12),
        Column("m.tg.before_kg", "9,4", 12),
        Column("m.tg.cut_name", "Striploin steak", 24),
        Column("meat.pieces", "17", 9),
        Column("m.tg.g_piece", "330", 12),
        Column("m.tg.trim", "No", 12),
        Column("m.tg.waste_kg", "0,6", 10),
    ]),
    Sheet("descongelado", "m.df.title", "m.df.sub", [
        Column("waste.serial", "8017-01", 14),
        Column("m.df.shift", "Cena", 10),
        Column("m.df.intake", "6", 12),
        Column("m.df.kg", "2,1", 12),
        Column("m.df.count", "2", 12),
        Column("m.df.left", "0,7", 12),
        Column("m.df.note", "", 22),
    ]),
    Sheet("inventario", "inv.title", "inv.sub", [
        Column("waste.serial", "8017-01", 14),
        Column("rec.component", "Striploin steak", 24),
        Column("meat.pieces", "15", 9),
        Column("inv.counted", "4,5", 12),
        Column("inv.expected", "4,5", 12),
        Column("inv.gap", "0", 10),
        Column("m.df.note", "", 22),
    ]),
    Sheet("maduracion", "m.ag.title", "m.ag.sub", [
        Column("m.rec.serial", "8017", 12),
        Column("m.rec.sku", "Ribeye AUS", 22),
        Column("m.ag.where", "Maduración", 14),
        Column("m.ag.days", "45", 8),
        Column("m.ag.weigh_kg", "7,6", 12),
        Column("m.ag.trim", "1,2", 12),
        Column("m.tg.trim", "0,3", 12),
        Column("m.ag.waste_kg", "0,9", 12),
        Column("m.df.note", "", 20),
    ]),
    Sheet("merma", "waste.title", "waste.sub", [
        Column("waste.lot", "TG-0010", 14),
        Column("waste.serial", "8017-01", 14),
        Column("rec.component", "Striploin steak", 24),
        Column("waste.kg", "0,6", 10),
        Column("waste.pieces", "2", 9),
        Column("waste.reason", "Caducado", 22),
    ]),
]

BY_CODE = {s.code: s for s in SHEETS}


def build(ws, sheet: Sheet, restaurant: Restaurant, lang: str = DEFAULT_LANG,
          blank_rows: int = BLANK_ROWS, site: str = ""):
    """Monta la hoja: cabecera, ejemplo en gris y rejilla en blanco.

    Con sede, la hoja lleva su nombre: colgada al lado de la balanza, el papel
    del obrador y el del local no se confunden.
    """
    casa = f"{restaurant.name} · {site}" if site else restaurant.name
    headers = ["Nº", t(lang, "common.date")] + [t(lang, c.key) for c in sheet.columns]
    last_col = len(headers)
    span = f"A1:{get_column_letter(last_col)}1"

    ws.sheet_view.rightToLeft = direction(lang) == "rtl"
    ws.sheet_view.showGridLines = False

    ws.merge_cells(span)
    title = ws["A1"]
    title.value = t(lang, sheet.title_key)
    title.font = Font(name=FONT, size=15, bold=True, color=INK)
    title.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 24

    ws.merge_cells(span.replace("1", "2"))
    subtitle = ws["A2"]
    subtitle.value = f"{casa} · {t(lang, sheet.sub_key)}"
    subtitle.font = Font(name=FONT, size=10, color=MUTED)

    row = 4
    col = 1
    for label in (t(lang, "sheet.restaurant"), t(lang, "sheet.shift"),
                  t(lang, "sheet.filled_by")):
        label_cell = ws.cell(row=row, column=col, value=f"{label}:")
        label_cell.font = Font(name=FONT, size=10, bold=True, color=INK)
        value_cell = ws.cell(row=row, column=col + 1,
                             value=casa if col == 1 else "")
        value_cell.font = Font(name=FONT, size=10, color=INK)
        value_cell.border = WRITE_LINE
        col += 2
    ws.row_dimensions[row].height = 20

    ws.merge_cells(f"A6:{get_column_letter(last_col)}6")
    legend = ws["A6"]
    legend.value = t(lang, "sheet.legend")
    legend.font = Font(name=FONT, size=9, italic=True, color=MUTED)
    ws.row_dimensions[6].height = 16

    head_row = 8
    for i, text in enumerate(headers, start=1):
        c = ws.cell(row=head_row, column=i, value=text)
        c.font = Font(name=FONT, size=10, bold=True, color=INK)
        c.fill = PatternFill("solid", fgColor=HEADER_FILL)
        c.border = GRID
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.row_dimensions[head_row].height = 34

    example_row = head_row + 1
    example = [t(lang, "sheet.example"), date.today().isoformat()] + \
              [c.example for c in sheet.columns]
    for i, value in enumerate(example, start=1):
        c = ws.cell(row=example_row, column=i, value=value)
        c.font = Font(name=FONT, size=9, italic=True, color=MUTED)
        c.fill = PatternFill("solid", fgColor=EXAMPLE_FILL)
        c.border = GRID
        c.alignment = Alignment(horizontal="center" if i <= 2 else "left")
    ws.row_dimensions[example_row].height = 18

    first_blank = example_row + 1
    for n in range(blank_rows):
        r = first_blank + n
        for i in range(1, last_col + 1):
            c = ws.cell(row=r, column=i)
            c.border = GRID
            c.font = Font(name=FONT, size=11, color=INK)
            if i == 1:
                c.value = n + 1
                c.alignment = Alignment(horizontal="center")
        ws.row_dimensions[r].height = 22

    sign_row = first_blank + blank_rows + 1
    sign = ws.cell(row=sign_row, column=1, value=f"{t(lang, 'sheet.signature')}:")
    sign.font = Font(name=FONT, size=10, bold=True, color=INK)
    for i in range(2, min(last_col, 6) + 1):
        ws.cell(row=sign_row, column=i).border = WRITE_LINE
    ws.row_dimensions[sign_row].height = 26

    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 13
    total_width = 19
    for i, column in enumerate(sheet.columns, start=3):
        ws.column_dimensions[get_column_letter(i)].width = column.width
        total_width += column.width

    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "portrait" if total_width <= PORTRAIT_MAX_WIDTH else "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"1:{head_row}"
    ws.print_area = f"A1:{get_column_letter(last_col)}{sign_row}"
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5
    ws.oddFooter.right.text = "&P / &N"
    ws.oddFooter.left.text = f"{casa} · {t(lang, sheet.title_key)}"
    ws.oddFooter.left.size = ws.oddFooter.right.size = 8
    ws.freeze_panes = ws.cell(row=head_row + 1, column=1)
    return ws


def workbook(code: str, restaurant: Restaurant, lang: str = DEFAULT_LANG,
             blank_rows: int = BLANK_ROWS, site: str = "") -> bytes:
    """Una hoja suelta, o el libro entero con `code="todo"`."""
    wanted = SHEETS if code == "todo" else [BY_CODE[code]]
    wb = Workbook()
    wb.remove(wb.active)
    used: set[str] = set()
    for sheet in wanted:
        ws = wb.create_sheet(sheet_title(t(lang, sheet.title_key), used))
        build(ws, sheet, restaurant, lang, blank_rows, site=site)
    stream = io.BytesIO()
    wb.save(stream)
    return stream.getvalue()


def filename(code: str, lang: str = DEFAULT_LANG) -> str:
    name = "todo" if code == "todo" else t(lang, BY_CODE[code].title_key)
    return filename_for(name, lang)
